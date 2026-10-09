from __future__ import annotations

from pathlib import Path
import shutil
import zipfile

import docx
import pytest
from docx.oxml.ns import qn

from docxwright import convert_file
from docxwright.bibliography import parse_bib, parse_names
from docxwright.converter import parse_document
from docxwright.model import Equation, Heading, Paragraph, Table
from docxwright.texmath import is_simple, parse_math, to_omml, to_text_runs

FIXTURES = Path(__file__).parent / "fixtures"


def plain(runs) -> str:
    return "".join(getattr(r, "text", "") for r in runs).replace("\u00a0", " ")


def test_document_structure():
    document, ctx, _ = parse_document(FIXTURES / "sample.tex")
    headings = [b for b in document.blocks if isinstance(b, Heading)]
    assert [h.number for h in headings if h.number] == ["1"]
    assert headings[-1].runs[0].text == "References"
    body = " ".join(plain(b.runs) for b in document.blocks if isinstance(b, Paragraph))
    assert "Delta et al. (2020) showed" in body
    assert "(see Delta et al. 2020; Rho 2019, p. 4)" in body
    assert "“quotes”, dashes – and — and an accent: Straße, Café." in body
    assert "Equation (1) in Section 1." in body
    assert "Table 1." in body
    equations = [b for b in document.blocks if isinstance(b, Equation)]
    assert [e.number for e in equations] == ["1", None]
    table = next(b for b in document.blocks if isinstance(b, Table))
    assert table.caption_above and table.rows[0][0].colspan == 2


def test_bibliography_formatting():
    document, _, _ = parse_document(FIXTURES / "sample.tex")
    refs = [plain(b.runs) for b in document.blocks if isinstance(b, Paragraph) and b.style == "Bibliography"]
    assert refs == [
        "Delta, Ann, Bob Sigma, and Cy Beta (May 2020). “A Study of Things”. In: Journal of Stuff 3.2, "
        "pp. 10–20. doi: 10.1000/xyz.",
        "Rho, Dee (2019). The Book. City: Pub. isbn: 123.",
    ]


def test_name_parsing():
    names = parse_names("Del Gamma, K. and Ann van Delta and {The Group}")
    assert (names[0].family, names[0].given) == ("Del Gamma", "K.")
    assert (names[1].family, names[1].prefix, names[1].given) == ("Delta", "van", "Ann")
    assert names[2].corporate


def test_bib_parser_handles_strings_and_concatenation():
    entries = parse_bib('@string{j = "Journal"}\n@article{k, title = {A {B} c}, journal = j # " X", year = 2001}')
    assert entries["k"].fields["journal"] == "Journal X"
    assert entries["k"].fields["title"] == "A {B} c"


def test_simple_math_becomes_text_runs():
    items = parse_math(r"p=0.05 \pm 1.2")
    assert is_simple(items)
    assert "".join(r.text for r in to_text_runs(items)) == "p = 0.05 ± 1.2"
    assert "".join(r.text for r in to_text_runs(parse_math(r"[\!-\!0.1,\,0.4]"))) == "[−0.1, 0.4]"


def test_complex_math_becomes_omml():
    items = parse_math(r"\frac{1}{S}\sum_{s=1}^{S} x_s")
    assert not is_simple(items)
    xml = to_omml(items, display=True)
    assert "<m:f>" in xml and "<m:nary>" in xml and 'm:val="∑"' in xml


def test_docx_output(tmp_path):
    out = tmp_path / "out.docx"
    convert_file(FIXTURES / "sample.tex", out, figures=False)
    document = docx.Document(str(out))
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == "A Small Test"
    assert "References" in texts
    with zipfile.ZipFile(out) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "<m:oMath" in xml


@pytest.mark.skipif(shutil.which("pdflatex") is None, reason="LaTeX not installed")
def test_figures_rendered_on_own_page(tmp_path):
    tex = tmp_path / "fig.tex"
    tex.write_text(
        "\\documentclass{article}\\usepackage{tikz}\\begin{document}\n"
        "Text before.\n\n\\begin{figure}\\centering\\begin{tikzpicture}\\draw (0,0) circle (1);"
        "\\end{tikzpicture}\\caption{A circle.}\\end{figure}\n\nText after.\n\\end{document}\n",
        encoding="utf-8",
    )
    out = tmp_path / "fig.docx"
    convert_file(tex, out)
    with zipfile.ZipFile(out) as archive:
        names = archive.namelist()
    assert any(name.startswith("word/media/") and name.endswith(".png") for name in names)
    document = docx.Document(str(out))
    paragraphs = [p for p in document.element.body if p.tag.endswith("}p")]
    texts = ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in paragraphs]
    picture = next(i for i, p in enumerate(paragraphs) if p.find(".//" + qn("w:drawing")) is not None)
    # the figure starts a new page right after the preceding text, the caption sits below it,
    # and the following text starts a new page again
    assert texts[picture - 1] == "Text before." and page_break_before(paragraphs[picture])
    assert texts[picture + 1] == "Figure 1: A circle."
    assert texts[picture + 2] == "Text after." and page_break_before(paragraphs[picture + 2])


def page_break_before(p) -> bool:
    ppr = p.find(qn("w:pPr"))
    return ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None


@pytest.mark.skipif(shutil.which("pdflatex") is None, reason="LaTeX not installed")
def test_consecutive_figures_single_break_outside_caption(tmp_path):
    figure = ("\\begin{figure}\\centering\\begin{tikzpicture}\\draw (0,0) circle (1);\\end{tikzpicture}"
              "\\caption{Caption %s.}\\end{figure}\n\n")
    tex = tmp_path / "figs.tex"
    tex.write_text("\\documentclass{article}\\usepackage{tikz}\\begin{document}\nText before.\n\n"
                   + figure % "one" + figure % "two" + "Text after.\n\\end{document}\n", encoding="utf-8")
    out = tmp_path / "figs.docx"
    convert_file(tex, out)
    document = docx.Document(str(out))
    body = document.element.body
    paragraphs = [p for p in body if p.tag.endswith("}p")]

    def text(p):
        return "".join(t.text or "" for t in p.iter(qn("w:t")))

    # No paragraphs that only hold a manual page break (they cause empty pages in some
    # viewers); figures start new pages through "page break before".
    assert not any(br.get(qn("w:type")) == "page" for p in paragraphs for br in p.iter(qn("w:br")))
    kinds = []
    for p in paragraphs:
        prefix = "|" if page_break_before(p) else ""
        if p.find(".//" + qn("w:drawing")) is not None:
            kinds.append(prefix + "figure")
        elif text(p):
            kinds.append(prefix + text(p))
    assert kinds == ["Text before.", "|figure", "Figure 1: Caption one.", "|figure",
                     "Figure 2: Caption two.", "|Text after."]
