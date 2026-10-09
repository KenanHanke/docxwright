"""Write the intermediate document model to a .docx file using python-docx.

When the model carries revision marks (see :mod:`docxwright.compare`), they are written
either as native Word tracked changes (``mode="tracked"``: insertions, deletions, formatting
changes, inserted/deleted paragraphs, table rows, equations and pictures, all of which can be
reviewed and accepted/rejected in Word or LibreOffice) or as highlighted insertions
(``mode="highlight"``: the new text with every insertion or modification highlighted, no
strike-through, changed figures framed).
"""

from __future__ import annotations

from copy import copy, deepcopy

from dataclasses import dataclass, field, replace
from datetime import datetime, timedelta, timezone
import math
from pathlib import Path
import difflib
import re

from docx import Document as new_document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

from .compare import diff_runs, recompute_vertical_merges, runs_with_rev
from .model import (
    ADDED,
    quoted_runs,
    whole_words,
    REMOVED,
    Break,
    Equation,
    Figure,
    CommentMark,
    Footnote,
    Heading,
    ListBlock,
    Math,
    PageBreak,
    Paragraph,
    Style,
    Table,
    TableOfContents,
    Text,
    TocEntry,
    figure_label,
    table_label,
    equation_label,
)

BODY_FONT = "Times New Roman"
MONO_FONT = "Courier New"
BASE_SIZE = 11.0
QUOTE_INDENT_PT = 25.0
FOOTNOTE_SIZE = 9.0
SIZE_PT = {"tiny": 6.0, "scriptsize": 8.0, "footnotesize": 9.0, "small": 10.0, "normal": 11.0,
           "large": 12.0, "Large": 14.4, "LARGE": 17.28, "huge": 20.74, "Huge": 24.88}
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 1.0
TEXT_WIDTH_IN = PAGE_WIDTH_IN - 2 * MARGIN_IN
TEXT_HEIGHT_IN = PAGE_HEIGHT_IN - 2 * MARGIN_IN
PARINDENT = Pt(16.5)  # 1.5em at 11pt, as in LaTeX
# A medium blue that stays readable on white and on dark (dark mode) page backgrounds.
LINK_COLOR = RGBColor(0x2F, 0x80, 0xD0)  # ~4:1 contrast against white and against a dark page
HIGHLIGHT_LINK_COLOR = "1F4E99"  # a link in highlighted text (~6:1 on the light green or blue)
RULE_COLOR = "808080"  # table rules: visible on white and on dark (dark mode) pages
COMMENT_TEXT_COLOR = RGBColor(0x00, 0x00, 0x00)  # LibreOffice up to 24.x draws comments on a light box, also in dark mode
LIST_INDENT_TWIPS = 360  # per nesting level


@dataclass
class ChangeOptions:
    """How revision marks are written."""

    mode: str = "tracked"  # "tracked" (native Word revisions) or "highlight"
    author: str = "docxwright"
    date: datetime = field(default_factory=lambda: datetime.now().astimezone())
    track_further_edits: bool = True  # switch on "Track Changes" for the revisers' own edits
    highlight_color: str = "C8F0C8"  # highlight mode: background of inserted/modified text
    frame_color: str = "1F9D3A"  # highlight mode: frame around new/changed figures
    move_color: str = "CFE2F3"  # highlight mode: background of moved text
    move_frame_color: str = "3D85C6"  # highlight mode: frame around moved equations
    comments: bool = True  # explanatory comments on changed figures
    baseline: str = ""  # the earlier version, e.g. '1a2b3c4 “Submitted version” (2 October 2026)'
    unchanged: bool = False  # the comparison found no visible change

    def legend(self) -> str:
        since = f"since {self.baseline}" if self.baseline else "since the earlier version"
        if self.unchanged:
            return f"There are no visible changes {since}: the text, figures, tables and equations are the same."
        if self.mode == "highlight":
            return (f"Changes {since} are highlighted: green background = new or reworded text, "
                    "blue background = moved text, dotted underline = formatting changed, green frame = "
                    "new or changed figure or equation. Deleted text is not shown: where a few words were "
                    "removed, the word next to the gap is green; comments quote longer removed passages, mark "
                    "removed paragraphs, figures, tables, table rows, equations and footnotes and added or "
                    "removed paragraph breaks, and explain moved text.")
        return (f"Changes {since} are shown as tracked changes, which can be accepted or rejected one by one "
                "or all at once (moved paragraphs as moves; a comment says where moved sentences came from). "
                "Formatting-only changes (for example a word no longer in a typewriter font) are tracked as "
                "formatting changes, which LibreOffice shows in bold. Comments explain changed figures and "
                "anything else that tracked changes cannot show.")

    @property
    def date_text(self) -> str:
        # Word writes (and displays) revision dates as local wall-clock time followed by "Z".
        return self.date.strftime("%Y-%m-%dT%H:%M:%SZ")


class DocxWriter:
    def __init__(self, changes: ChangeOptions | None = None):
        self.doc = new_document()
        self.changes = changes
        self.revision_id = 0
        self.toc_entries: list[tuple[int, list, str]] = []  # (level, runs, bookmark)
        # The next block starts on a new page (after a figure, \newpage) in these versions
        # ("old"/"new"; a break of only one version is a tracked paragraph property change).
        self.pending_break: set[str] = set()
        self.pending_notes: list[str] = []  # comments for the next block (highlight mode: removed items)
        # Tracked mode: comments for the next unchanged paragraph (a comment inside a tracked
        # insertion keeps LibreOffice from rejecting it).
        self.pending_tracked_notes: list[str] = []
        self.pending_break_notes: list[str] = []  # tracked mode: changed paragraph breaks, for the next paragraph
        self.list_nums: dict[tuple[str, int], int] = {}
        self.footnotes = None  # <w:footnotes> element of the footnotes part, created on demand
        self.footnote_part = None
        self.footnote_id = 0
        self.setup_page()
        self.setup_styles()
        self.setup_numbering()

    @property
    def tracked(self) -> bool:
        return self.changes is not None and self.changes.mode == "tracked"

    @property
    def highlight(self) -> bool:
        return self.changes is not None and self.changes.mode == "highlight"

    # ------------------------------------------------------------------ setup

    def setup_page(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Inches(PAGE_WIDTH_IN)
        section.page_height = Inches(PAGE_HEIGHT_IN)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Inches(MARGIN_IN))
        # Page numbers centered in the footer, as in the LaTeX article class.
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.add_field(footer, "PAGE", "1")
        settings = self.doc.settings.element
        compat = settings.find(qn("w:compat"))
        if compat is None:
            compat = OxmlElement("w:compat")
            settings.append(compat)
        # Lines ending in a manual line break are not stretched in justified paragraphs.
        if compat.find(qn("w:doNotExpandShiftReturn")) is None:
            compat.insert(0, OxmlElement("w:doNotExpandShiftReturn"))
        if self.changes is not None and self.changes.track_further_edits:
            insert_setting(settings, OxmlElement("w:trackRevisions"))

    def setup_styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        set_font(normal, BODY_FONT, BASE_SIZE)
        pf = normal.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.1
        pf.widow_control = True
        # Document defaults: automatic text color (white on dark backgrounds).
        defaults = self.doc.styles.element.find(qn("w:docDefaults"))
        if defaults is not None:
            for rpr in defaults.iter(qn("w:rPr")):
                set_auto_color(rpr)

        heading_sizes = {1: 14.4, 2: 12.0, 3: 11.0, 4: 11.0, 5: 11.0}
        heading_space = {1: (18, 9), 2: (14, 6), 3: (12, 6), 4: (10, 4), 5: (8, 3)}
        for level, size in heading_sizes.items():
            style = styles[f"Heading {level}"]
            set_font(style, BODY_FONT, size, bold=True, italic=False)
            style.paragraph_format.space_before = Pt(heading_space[level][0])
            style.paragraph_format.space_after = Pt(heading_space[level][1])
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

        title = styles["Title"]
        set_font(title, BODY_FONT, 17.28, bold=False)
        remove_child(title.element.pPr, "w:pBdr")
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.space_after = Pt(12)
        title.paragraph_format.line_spacing = 1.1
        rpr = title.element.get_or_add_rPr()
        for tag in ("w:spacing", "w:kern"):
            remove_child(rpr, tag)

        author = self.paragraph_style("Author", base="Normal", size=12.0)
        subtitle = self.paragraph_style("Subtitle", base="Normal", size=14.4)
        subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.first_line_indent = Pt(0)
        subtitle.paragraph_format.space_after = Pt(6)
        author.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author.paragraph_format.space_after = Pt(4)
        affil = self.paragraph_style("Affiliation", base="Normal", size=10.0)
        affil.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        affil.paragraph_format.left_indent = Pt(14.3)
        affil.paragraph_format.first_line_indent = Pt(-14.3)

        caption = styles["Caption"]
        set_font(caption, BODY_FONT, BASE_SIZE, bold=False, italic=False)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        caption.paragraph_format.space_before = Pt(10)
        caption.paragraph_format.space_after = Pt(10)

        bib = self.paragraph_style("Bibliography", base="Normal", size=BASE_SIZE)
        bib.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        bib.paragraph_format.left_indent = Pt(11)
        bib.paragraph_format.first_line_indent = Pt(-11)
        bib.paragraph_format.space_after = Pt(3)

        fn_text = self.paragraph_style("Footnote Text", base="Normal", size=FOOTNOTE_SIZE, style_id="FootnoteText")
        fn_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fn_text.paragraph_format.space_after = Pt(1)
        if "Footnote Reference" not in [s.name for s in styles]:
            fn_ref = styles.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
            fn_ref.element.set(qn("w:styleId"), "FootnoteReference")
            fn_ref.font.superscript = True
            fn_ref.priority = 99

        # Comments (explanations of changed figures, moved paragraphs): Word's built-in styles.
        comment_text = self.paragraph_style("annotation text", base="Normal", size=10.0, style_id="CommentText")
        comment_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        comment_text.paragraph_format.first_line_indent = Pt(0)
        comment_text.paragraph_format.space_after = Pt(0)
        comment_text.priority = 99
        if "annotation reference" not in [s.name for s in styles]:
            comment_ref = styles.add_style("annotation reference", WD_STYLE_TYPE.CHARACTER)
            comment_ref.element.set(qn("w:styleId"), "CommentReference")
            comment_ref.font.size = Pt(8)
            comment_ref.priority = 99

        labeled = self.paragraph_style("Bibliography Labeled", base="Bibliography", size=BASE_SIZE)
        labeled.paragraph_format.left_indent = Pt(30)
        labeled.paragraph_format.first_line_indent = Pt(-30)

        toc_heading = styles["TOC Heading"]
        set_font(toc_heading, BODY_FONT, 14.4, bold=True, italic=False)
        toc_heading.paragraph_format.space_before = Pt(18)
        toc_heading.paragraph_format.space_after = Pt(9)
        for level, indent in ((1, 0), (2, 16.5), (3, 41.8)):
            style = self.paragraph_style(f"toc {level}", base="Normal", size=BASE_SIZE, style_id=f"TOC{level}")
            style.paragraph_format.left_indent = Pt(indent)
            style.paragraph_format.space_before = Pt(8 if level == 1 else 0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if level == 1:
                style.font.bold = True
            style.priority = 39

        equation = self.paragraph_style("Equation", base="Normal", size=BASE_SIZE)
        equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        equation.paragraph_format.space_before = Pt(6)
        equation.paragraph_format.space_after = Pt(6)

        figure = self.paragraph_style("Figure", base="Normal", size=BASE_SIZE)
        figure.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure.paragraph_format.keep_with_next = True
        figure.paragraph_format.space_after = Pt(4)

        listp = self.paragraph_style("List Paragraph", base="Normal", size=None)
        listp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        listp.paragraph_format.space_after = Pt(2)

        if "Hyperlink" not in [s.name for s in styles]:
            link = styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
            link.hidden = False
            link.priority = 99
        link = styles["Hyperlink"]
        link.font.color.rgb = LINK_COLOR
        link.font.underline = True

        for name in ("Subtitle", "Quote", "Intense Quote", "List Bullet", "List Number", "List Continue",
                     "No Spacing", "Body Text", "Header", "Footer"):
            if name in [s.name for s in styles]:
                set_font(styles[name], BODY_FONT if name not in ("Header", "Footer") else None, None)
        # Remove theme colors everywhere: in dark mode, explicit (non-automatic) dark text
        # colors stay dark on the dark page background and become unreadable.
        for style in styles.element.iter(qn("w:style")):
            for color in style.iter(qn("w:color")):
                if color.get(qn("w:val")) in (None, "000000") or color.get(qn("w:themeColor")):
                    if style.get(qn("w:styleId")) != "Hyperlink":
                        make_auto(color)

    def paragraph_style(self, name: str, base: str, size: float | None, style_id: str | None = None):
        styles = self.doc.styles
        existing = [s for s in styles if s.name == name]
        style = existing[0] if existing else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if style_id:
            style.element.set(qn("w:styleId"), style_id)
        style.base_style = styles[base]
        style.quick_style = True
        if size is not None:
            style.font.size = Pt(size)
        return style

    def setup_numbering(self) -> None:
        """Abstract numbering definitions for LaTeX-like bullet and enumerate lists."""
        numbering = self.doc.part.numbering_part.element
        self.numbering = numbering
        self.abstract_ids = {}
        bullet_chars = ["•", "–", "∗", "·", "•", "–", "∗", "·", "•"]
        number_formats = [("decimal", "%1."), ("lowerLetter", "(%2)"), ("lowerRoman", "%3."), ("upperLetter", "%4."),
                          ("decimal", "%5."), ("lowerLetter", "%6."), ("lowerRoman", "%7."), ("decimal", "%8."),
                          ("lowerLetter", "%9.")]
        existing = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        next_id = max(existing, default=0) + 1
        for kind in ("bullet", "number"):
            abstract = OxmlElement("w:abstractNum")
            abstract.set(qn("w:abstractNumId"), str(next_id))
            multi = OxmlElement("w:multiLevelType")
            multi.set(qn("w:val"), "hybridMultilevel")
            abstract.append(multi)
            for level in range(9):
                lvl = OxmlElement("w:lvl")
                lvl.set(qn("w:ilvl"), str(level))
                start = OxmlElement("w:start")
                start.set(qn("w:val"), "1")
                fmt = OxmlElement("w:numFmt")
                text = OxmlElement("w:lvlText")
                if kind == "bullet":
                    fmt.set(qn("w:val"), "bullet")
                    text.set(qn("w:val"), bullet_chars[level])
                else:
                    fmt.set(qn("w:val"), number_formats[level][0])
                    text.set(qn("w:val"), number_formats[level][1])
                jc = OxmlElement("w:lvlJc")
                jc.set(qn("w:val"), "left")
                ppr = OxmlElement("w:pPr")
                ind = OxmlElement("w:ind")
                ind.set(qn("w:left"), str(LIST_INDENT_TWIPS * (level + 1)))
                ind.set(qn("w:hanging"), str(LIST_INDENT_TWIPS))
                ppr.append(ind)
                for child in (start, fmt, text, jc, ppr):
                    lvl.append(child)
                if kind == "bullet":
                    rpr = OxmlElement("w:rPr")
                    fonts = OxmlElement("w:rFonts")
                    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                        fonts.set(qn(attr), BODY_FONT)
                    rpr.append(fonts)
                    lvl.append(rpr)
                abstract.append(lvl)
            # abstractNum elements must precede num elements.
            first_num = numbering.find(qn("w:num"))
            if first_num is not None:
                first_num.addprevious(abstract)
            else:
                numbering.append(abstract)
            self.abstract_ids[kind] = next_id
            next_id += 1

    def list_num_id(self, kind: str, list_id: int, level: int = 0, custom: tuple | None = None) -> int:
        key = (kind, list_id if (kind == "number" or custom) else 0, custom)
        if key in self.list_nums:
            return self.list_nums[key]
        existing = [int(n.get(qn("w:numId"))) for n in self.numbering.findall(qn("w:num"))]
        num_id = max(existing, default=0) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), str(self.abstract_ids[kind]))
        num.append(abstract)
        if kind == "number" or custom:
            for lvl_index in range(9):
                override = OxmlElement("w:lvlOverride")
                override.set(qn("w:ilvl"), str(lvl_index))
                start = OxmlElement("w:startOverride")
                start.set(qn("w:val"), str(custom[2] if custom and lvl_index == level else 1))
                override.append(start)
                if custom and lvl_index == level:
                    override.append(self.custom_level(lvl_index, custom))
                num.append(override)
        self.numbering.append(num)
        self.list_nums[key] = num_id
        return num_id

    def custom_level(self, level: int, custom: tuple):
        fmt_name, text, start_value = custom[:3]
        emphasis = custom[3] if len(custom) > 3 else ""
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), str(start_value))
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), fmt_name)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        width = max(LIST_INDENT_TWIPS, 120 * (len(text) + 1))
        ind.set(qn("w:left"), str(LIST_INDENT_TWIPS * level + width))
        ind.set(qn("w:hanging"), str(width))
        ppr.append(ind)
        for child in (start, fmt, lvl_text, jc, ppr):
            lvl.append(child)
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), BODY_FONT)
        rpr.append(fonts)
        if "b" in emphasis:
            rpr.append(OxmlElement("w:b"))
        if "i" in emphasis:
            rpr.append(OxmlElement("w:i"))
        lvl.append(rpr)
        return lvl

    # ------------------------------------------------------------------ revisions

    def next_revision_id(self) -> str:
        self.revision_id += 1
        return str(self.revision_id)

    def revision_element(self, tag: str):
        element = OxmlElement(tag)
        element.set(qn("w:id"), self.next_revision_id())
        element.set(qn("w:author"), self.changes.author)
        element.set(qn("w:date"), self.changes.date_text)
        return element

    def place(self, container, element, rev: str | None) -> None:
        """Append a run-level element to a paragraph/hyperlink, wrapped in w:ins/w:del if needed."""
        if self.tracked and rev in REVISION_TAGS:
            tag = REVISION_TAGS[rev]
            last = container[-1] if len(container) else None
            if last is not None and last.tag == qn(tag) and last.get("_merge") == "1":
                wrapper = last
            else:
                wrapper = self.revision_element(tag)
                wrapper.set("_merge", "1")
                container.append(wrapper)
            if rev in REMOVED:
                to_deleted(element)
            wrapper.append(element)
        else:
            container.append(element)

    def mark_paragraph(self, paragraph, rev: str | None) -> None:
        """Revision mark on the paragraph mark (inserted / deleted paragraph break)."""
        if not self.tracked or rev not in REVISION_TAGS:
            return
        ppr = paragraph._p.get_or_add_pPr()
        rpr = ppr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ppr.append(rpr)
            reorder_children(ppr, PPR_ORDER)
        rpr.insert(0, self.revision_element(REVISION_TAGS[rev]))

    def style_change(self, paragraph, old_style: str) -> None:
        """Tracked mode: the paragraph had style OLD_STYLE in the old version."""
        if not self.tracked or old_style not in [s.name for s in self.doc.styles]:
            return
        ppr = paragraph._p.get_or_add_pPr()
        old_ppr = OxmlElement("w:pPr")
        for child in ppr:
            if child.tag not in (qn("w:rPr"), qn("w:sectPr"), qn("w:pPrChange")):
                old_ppr.append(deepcopy(child))
        style = old_ppr.find(qn("w:pStyle"))
        if style is None:
            style = OxmlElement("w:pStyle")
            old_ppr.insert(0, style)
        style.set(qn("w:val"), self.doc.styles[old_style].style_id)
        change = self.revision_element("w:pPrChange")
        change.append(old_ppr)
        ppr.append(change)
        reorder_children(ppr, PPR_ORDER)

    def list_change(self, paragraph, block: Paragraph) -> None:
        """Tracked mode: the list item had another kind of label in the old version (an
        automatic number that became \\item[RQ1], an itemize that became an enumerate, ...)."""
        if not self.tracked:
            return
        old_kind, old_format = block.old_list
        scratch = self.doc.add_paragraph(style="List Paragraph")
        self.list_format(scratch, replace(block, list_kind=old_kind, list_format=old_format, old_list=None))
        scratch._p.getparent().remove(scratch._p)
        layout = (qn("w:numPr"), qn("w:ind"), qn("w:jc"), qn("w:tabs"))
        ppr = paragraph._p.get_or_add_pPr()
        old_ppr = OxmlElement("w:pPr")
        for child in ppr:
            if child.tag not in layout + (qn("w:rPr"), qn("w:sectPr"), qn("w:pPrChange")):
                old_ppr.append(deepcopy(child))
        scratch_ppr = scratch._p.find(qn("w:pPr"))
        for child in scratch_ppr if scratch_ppr is not None else []:
            if child.tag in layout:
                old_ppr.append(deepcopy(child))
        reorder_children(old_ppr, PPR_ORDER)
        change = self.revision_element("w:pPrChange")
        change.append(old_ppr)
        ppr.append(change)
        reorder_children(ppr, PPR_ORDER)

    def mark_move(self, paragraph, block) -> None:
        """Tracked mode: the range markers that tie the old and new position of a moved paragraph."""
        if not self.tracked or not block.move or block.mark_rev not in ("movefrom", "moveto"):
            return
        kind = "moveFrom" if block.mark_rev == "movefrom" else "moveTo"
        start = self.revision_element(f"w:{kind}RangeStart")
        start.set(qn("w:name"), block.move)
        end = OxmlElement(f"w:{kind}RangeEnd")
        end.set(qn("w:id"), start.get(qn("w:id")))
        p = paragraph._p
        ppr = p.find(qn("w:pPr"))
        if ppr is not None:
            ppr.addnext(start)
        else:
            p.insert(0, start)
        p.append(end)

    def mark_row(self, row, rev: str | None) -> None:
        if not self.tracked or rev not in ("ins", "del"):
            return
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(self.revision_element("w:ins" if rev == "ins" else "w:del"))

    def apply_highlight(self, run_element, keep_color: bool = False, fill: str | None = None,
                        text_color: str = "1A1A1A") -> None:
        rpr = run_element.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            run_element.insert(0, rpr)
        if not keep_color:
            # Explicit dark text on the light highlight: "automatic" text turns white in dark
            # mode, which would be unreadable on the light background.
            remove_child(rpr, "w:color")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), text_color)
            rpr.append(color)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill or self.changes.highlight_color)
        remove_child(rpr, "w:shd")
        rpr.append(shd)
        reorder_children(rpr, RPR_ORDER)

    # ------------------------------------------------------------------ runs

    def add_runs(self, paragraph, runs: list, base_size: float | None = None, force_rev: str | None = None,
                 math_size: float | None = None) -> None:
        container = paragraph._p
        current_link = None
        current_url = None
        for run in runs:
            # (A forced mark - the contents of a footnote that moved with its reference - does
            # not override the text's own changes.)
            own = getattr(run, "rev", None)
            rev = own if force_rev and own in ("ins", "del", "fmt") else (force_rev or own)
            if self.highlight and rev in REMOVED:
                continue
            if isinstance(run, Text):
                if not run.text:
                    continue
                if run.style.url:
                    if current_link is None or current_url != run.style.url or current_link is not container[-1]:
                        current_link = self.hyperlink_element(paragraph, run.style.url)
                        current_url = run.style.url
                        container.append(current_link)
                    target = current_link
                    style = Style(bold=run.style.bold, italic=run.style.italic, mono=run.style.mono,
                                  smallcaps=run.style.smallcaps, vert=run.style.vert)
                    r = self.make_run(run.text, style, base_size)
                    rstyle = OxmlElement("w:rStyle")
                    rstyle.set(qn("w:val"), self.doc.styles["Hyperlink"].style_id)
                    rpr = r.find(qn("w:rPr"))
                    if rpr is None:
                        rpr = OxmlElement("w:rPr")
                        r.insert(0, rpr)
                    rpr.insert(0, rstyle)
                else:
                    target = container
                    r = self.make_run(run.text, run.style, base_size)
                    current_link = None
                # (Text moved and reformatted is marked both ways.)
                reformatted = rev == "fmt" or (rev == "moveto" and run.old_style is not None
                                               and run.old_style != run.style)
                if reformatted:
                    if self.tracked and run.old_style is not None:
                        self.format_change(r, run.old_style, base_size)
                    elif self.highlight:
                        mark_format_change(r)
                document_language = getattr(self, "language", None)
                if run.lang in LANGUAGE_TAGS and run.lang != document_language:
                    # (Proofing language of a phrase in another language: \foreignlanguage, a footnote.)
                    set_run_language(r, LANGUAGE_TAGS[run.lang])
                elif run.lang is None and getattr(self, "paragraph_language", None) not in (None, document_language) \
                        and document_language in LANGUAGE_TAGS:
                    # (Text in the document's language within a paragraph mostly in another one.)
                    set_run_language(r, LANGUAGE_TAGS[document_language])
                if rev in ADDED and self.highlight:
                    moved = rev == "moveto" or getattr(run, "moved", False)
                    # (A link keeps its blue, darkened for contrast on the light background.)
                    self.apply_highlight(r, fill=self.changes.move_color if moved else None,
                                         text_color=HIGHLIGHT_LINK_COLOR if run.style.url else "1A1A1A")
                self.place(target, r, rev)
            elif isinstance(run, Math):
                current_link = None
                marked = self.highlight and rev in ADDED
                fill = self.changes.move_color if rev == "moveto" else None
                if marked:
                    # LibreOffice ignores the shading inside formulas: shaded thin spaces on both
                    # sides mark a new or changed inline formula there too.
                    edge = self.make_run("\u2009", Style(), base_size)
                    self.apply_highlight(edge, fill=fill)
                    self.place(container, edge, rev)
                self.place_math(container, run.omml, rev, display=False, size=math_size or base_size)
                if marked:
                    edge = self.make_run("\u2009", Style(), base_size)
                    self.apply_highlight(edge, fill=fill)
                    self.place(container, edge, rev)
            elif isinstance(run, CommentMark):
                current_link = None
                anchor = OxmlElement("w:r")
                container.append(anchor)
                self.comment(paragraph, run.text, anchor=anchor)
            elif isinstance(run, Footnote):
                current_link = None
                r = self.footnote_reference(run, rev)
                marked = self.highlight and rev in ADDED
                fill = self.changes.move_color if rev == "moveto" or run.moved else None
                if marked:
                    # (Its own text colour: LibreOffice paints no shading behind a footnote mark,
                    # so dark text would vanish on a dark page. Shaded thin spaces on both sides
                    # mark the new footnote there.)
                    self.apply_highlight(r, keep_color=True, fill=fill)
                    edge = self.make_run("\u2009", Style(), base_size)
                    self.apply_highlight(edge, fill=fill)
                    self.place(container, edge, rev)
                self.place(container, r, rev)
                if marked:
                    edge = self.make_run("\u2009", Style(), base_size)
                    self.apply_highlight(edge, fill=fill)
                    self.place(container, edge, rev)
                if run.mark is not None:
                    # An invisible character of its own after a custom mark: LibreOffice would
                    # otherwise take the mark of a following footnote as this footnote's label.
                    self.place(container, self.make_run("\u200c", Style(), None), rev)
            elif isinstance(run, Break):
                current_link = None
                r = OxmlElement("w:r")
                br = OxmlElement("w:br")
                if run.kind == "page":
                    br.set(qn("w:type"), "page")
                r.append(br)
                self.place(container, r, rev)
        for wrapper in container.iter(*(qn(tag) for tag in REVISION_TAGS.values())):
            if "_merge" in wrapper.attrib:
                del wrapper.attrib["_merge"]

    def footnote_reference(self, footnote: Footnote, rev: str | None):
        """Create the footnote (content in the footnotes part) and return its reference run."""
        from docx.opc.constants import CONTENT_TYPE as CT
        from docx.opc.packuri import PackURI
        from docx.opc.part import XmlPart
        from docx.text.paragraph import Paragraph as DocxParagraph

        if self.footnotes is None:
            xml = ('<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                   'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                   '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
                   'w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
                   '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" '
                   'w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
                   '</w:footnotes>')
            element = parse_xml(xml)
            part = XmlPart(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, element, self.doc.part.package)
            self.doc.part.relate_to(part, RT.FOOTNOTES)
            self.footnotes = element
            self.footnote_part = part
            settings = self.doc.settings.element
            pr = OxmlElement("w:footnotePr")
            for fid in ("-1", "0"):
                ref = OxmlElement("w:footnote")
                ref.set(qn("w:id"), fid)
                pr.append(ref)
            compat = settings.find(qn("w:compat"))
            if compat is not None:
                compat.addprevious(pr)
            else:
                settings.append(pr)
        self.footnote_id += 1
        fid = str(self.footnote_id)
        note = OxmlElement("w:footnote")
        note.set(qn("w:id"), fid)
        p = OxmlElement("w:p")
        note.append(p)
        self.footnotes.append(note)

        class _Parent:
            part = self.footnote_part

        paragraph = DocxParagraph(p, _Parent())
        ppr = p.get_or_add_pPr()
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), "FootnoteText")
        ppr.insert(0, pstyle)
        mark = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rstyle = OxmlElement("w:rStyle")
        rstyle.set(qn("w:val"), "FootnoteReference")
        rpr.append(rstyle)
        mark.append(rpr)
        if footnote.mark is not None:
            add_text_elements(mark, footnote.mark)
            # LibreOffice drops the first run of the first footnote (where it expects the
            # automatic mark); an empty run keeps the custom mark visible.
            self.place(p, OxmlElement("w:r"), rev)
        else:
            mark.append(OxmlElement("w:footnoteRef"))
        self.place(p, mark, rev)
        if len(p) and p[-1].get("_merge") == "1":
            # The text gets a wrapper of its own: LibreOffice marks only the number otherwise.
            p[-1].set("_merge", "0")
        space = self.make_run(" ", Style(), None)
        self.place(p, space, rev)
        # The contents of a moved footnote follow its reference (the move itself is marked there).
        # (The contents of a footnote that moved with its reference keep only their own changes:
        # the footnote goes with its reference when the move is accepted or rejected.)
        inner = {"movefrom": None, "moveto": "moveto" if self.highlight else None}.get(rev, rev)
        if self.highlight and footnote.moved and rev == "ins":
            inner = "moveto"  # (a footnote that moved to another sentence: blue)
        self.add_runs(paragraph, footnote.runs, force_rev=inner if inner in REVISION_TAGS else None,
                      math_size=FOOTNOTE_SIZE)
        self.mark_paragraph(paragraph, rev)
        ref = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rstyle = OxmlElement("w:rStyle")
        rstyle.set(qn("w:val"), "FootnoteReference")
        rpr.append(rstyle)
        ref.append(rpr)
        reference = OxmlElement("w:footnoteReference")
        if footnote.mark is not None:
            # A custom mark: Word and LibreOffice disagree on whether the mark is repeated in the
            # footnote area, so the reference itself gets an invisible custom mark (a zero-width
            # space, which also keeps it out of the automatic numbering), the visible mark is
            # ordinary superscript text before it and the footnote starts with the mark as text.
            if not footnote.hidden:
                add_text_elements(ref, footnote.mark)
            reference.set(qn("w:customMarkFollows"), "1")
        reference.set(qn("w:id"), fid)
        ref.append(reference)
        if footnote.mark is not None:
            add_text_elements(ref, "\u200b")
        return ref

    def make_run(self, text: str, style: Style, size: float | None):
        r = OxmlElement("w:r")
        rpr = run_properties(style, size)
        if len(rpr):
            r.append(rpr)
        add_text_elements(r, text)
        return r

    def format_change(self, r, old_style: Style, size: float | None) -> None:
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        change = self.revision_element("w:rPrChange")
        old = run_properties(old_style, size)
        change.append(old)
        rpr.append(change)

    def place_math(self, container, omml: str, rev: str | None, display: bool, justify: str | None = None,
                   size: float | None = None) -> None:
        element = parse_xml(omml)
        if size:
            # (LibreOffice sets formulas in its default size otherwise, e.g. in footnotes.)
            for mr in element.iter(qn("m:r")):
                rpr = mr.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    mrpr = mr.find(qn("m:rPr"))
                    if mrpr is not None:
                        mrpr.addnext(rpr)
                    else:
                        mr.insert(0, rpr)
                remove_child(rpr, "w:sz")
                sz = OxmlElement("w:sz")
                sz.set(qn("w:val"), str(int(round(size * 2))))
                rpr.append(sz)
                reorder_children(rpr, RPR_ORDER)
        if rev in ("ins", "fmt", "moveto") and self.highlight:
            for mr in element.iter(qn("m:r")):
                rpr = mr.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    mrpr = mr.find(qn("m:rPr"))
                    if mrpr is not None:
                        mrpr.addnext(rpr)
                    else:
                        mr.insert(0, rpr)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), self.changes.move_color if rev == "moveto" else self.changes.highlight_color)
                remove_child(rpr, "w:color")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "1A1A1A")
                rpr.append(color)
                rpr.append(shd)
                reorder_children(rpr, RPR_ORDER)
        if display:
            para = OxmlElement("m:oMathPara")
            if justify:
                para_pr = OxmlElement("m:oMathParaPr")
                jc = OxmlElement("m:jc")
                jc.set(qn("m:val"), justify)
                para_pr.append(jc)
                para.append(para_pr)
            para.append(element)
            element = para
        if rev in REVISION_TAGS and self.tracked:
            # The whole equation object is inserted/deleted (accepting or rejecting the change
            # then leaves no empty equation zone behind).
            wrapper = self.revision_element(REVISION_TAGS[rev])
            wrapper.append(element)
            container.append(wrapper)
        else:
            container.append(element)

    def hyperlink_element(self, paragraph, url: str):
        rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel_id)
        link.set(qn("w:history"), "1")
        return link

    def add_field(self, paragraph, instruction: str, cached: str) -> None:
        def fld(kind):
            r = OxmlElement("w:r")
            char = OxmlElement("w:fldChar")
            char.set(qn("w:fldCharType"), kind)
            r.append(char)
            return r

        paragraph._p.append(fld("begin"))
        r = OxmlElement("w:r")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f" {instruction} "
        r.append(instr)
        paragraph._p.append(r)
        paragraph._p.append(fld("separate"))
        paragraph.add_run(cached)
        paragraph._p.append(fld("end"))

    def add_bookmark(self, paragraph, name: str) -> None:
        # (Annotation ids - bookmarks, revisions, comments - share one id space.)
        bookmark_id = self.next_revision_id()
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), bookmark_id)
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), bookmark_id)
        ppr = paragraph._p.find(qn("w:pPr"))
        index = 1 if ppr is not None else 0
        paragraph._p.insert(index, start)
        paragraph._p.append(end)

    # ------------------------------------------------------------------ blocks

    def write(self, blocks: list) -> None:
        blocks = list(blocks)
        if self.highlight:
            blocks = accept_deletions(blocks)
        # Collect TOC entries up front (they are written where \tableofcontents occurs).
        for index, block in enumerate(blocks):
            if isinstance(block, Heading):
                if not block.bookmark:
                    block.bookmark = f"_Toc{100000 + index}"
                if block.toc and not heading_fully_deleted(block):
                    self.toc_entries.append(TocEntry(block.level, heading_runs(block, label=False), block.bookmark))
        self.fit_bibliography_labels(blocks)
        index = 0
        while index < len(blocks):
            block = blocks[index]
            if isinstance(block, Equation) and block.group is not None and block.parts is not None:
                group = [block]
                while (index + 1 < len(blocks) and isinstance(blocks[index + 1], Equation)
                       and blocks[index + 1].group == block.group and blocks[index + 1].parts is not None):
                    index += 1
                    group.append(blocks[index])
                self.write_equation_group(group)
                index += 1
                continue
            nxt = blocks[index + 1] if index + 1 < len(blocks) else None
            index += 1
            if isinstance(block, RemovalNote):
                self.pending_notes.append(block.text)
                continue
            if isinstance(block, PageBreak):
                self.pending_break |= versions_of(block.rev)
                continue
            body = self.doc.element.body
            before = len(content_children(body))
            if isinstance(block, Paragraph) and block.style == "Caption" and isinstance(nxt, Figure) \
                    and not self.pending_break and nxt.rev != "del":
                # (Text written above side-by-side figures - a shared title or source line - goes on
                # the figure's page: the page break comes before it, not before the picture.)
                pending_notes, self.pending_notes = self.pending_notes, []
                self.write_block(block, nxt)
                self.after_block(body, before, set(), pending_notes)
                written = content_children(body)[before:]
                if written and written[0].tag == qn("w:p"):
                    self.page_break_before(written[0], {"old", "new"})
                    written[-1].get_or_add_pPr().keepNext_val = True
                    self.figure_break_done = True
                continue
            if isinstance(block, Paragraph) and block.style == "Caption" and self.pending_break:
                # (The caption below side-by-side figures stays on the page of the figure above;
                # the page break comes after it.)
                deferred, self.pending_break = self.pending_break, set()
                pending_notes, self.pending_notes = self.pending_notes, []
                self.write_block(block, nxt)
                self.after_block(body, before, set(), pending_notes)
                self.pending_break |= deferred
                continue
            pending_break, self.pending_break = self.pending_break, set()
            pending_notes, self.pending_notes = self.pending_notes, []
            self.write_block(block, nxt)
            self.after_block(body, before, pending_break, pending_notes)
            if self.tracked and paragraph_break_changed(block, nxt):
                # (Only the paragraph mark changed, which is easy to miss: a comment where it was -
                # at the start of the next paragraph, which has unchanged text of its own.)
                self.pending_break_notes.append("A paragraph break was added here." if block.mark_rev == "ins"
                                                else "Two paragraphs were joined here.")
        if self.pending_notes and self.doc.paragraphs:
            for note in self.pending_notes:
                self.comment(self.doc.paragraphs[-1], note)
        if self.changes is not None and self.doc.paragraphs:
            # What the markings mean and what they compare with, at the start of the document.
            # (Not in a tracked paragraph: LibreOffice's Reject All mishandles comments there.)
            untracked = [p for p in self.doc.paragraphs if not any(
                p._p.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn(tag)}") is not None for tag in REVISION_TAGS.values())]
            first = next((p for p in untracked if p.style is not None and p.style.name == "Title"),
                         untracked[0] if untracked else None)
            if first is not None:
                self.comment(first, self.changes.legend(), at_start=True)
        # A document always ends with a paragraph that cannot be removed: if the last one is a
        # tracked insertion/deletion, an untracked empty paragraph keeps accepting/rejecting from
        # leaving a stray empty paragraph with that paragraph's style.
        children = content_children(self.doc.element.body)
        last = children[-1] if children else None
        if last is not None and (last.tag == qn("w:tbl") or any(
                last.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn(tag)}") is not None for tag in REVISION_TAGS.values())):
            self.doc.add_paragraph(style="Normal").paragraph_format.space_after = Pt(0)

    def after_block(self, body, before: int, pending_break: set, pending_notes: list) -> None:
        """Apply a pending page break and pending notes to the content just written."""
        new = [child for child in content_children(body)[before:] if child.tag in (qn("w:p"), qn("w:tbl"))]
        if self.pending_break_notes and new and new[0].tag == qn("w:p"):
            from docx.text.paragraph import Paragraph as DocxParagraph

            for note in self.pending_break_notes:
                self.comment(DocxParagraph(new[0], self.doc._body), note, at_start=True)
            self.pending_break_notes = []
        if self.pending_tracked_notes and new and new[0].tag == qn("w:p") and _mark_of(new[0]) is None and any(
                t.getparent() is not None and t.getparent().getparent() is new[0] for t in new[0].iter(qn("w:t"))):
            from docx.text.paragraph import Paragraph as DocxParagraph

            for note in self.pending_tracked_notes:
                self.comment(DocxParagraph(new[0], self.doc._body), note, at_start=True)
            self.pending_tracked_notes = []
        if not new:
            self.pending_break = self.pending_break | pending_break
            self.pending_notes = pending_notes + self.pending_notes
            return
        first = new[0]
        if not self.tracked and pending_break:
            pending_break = {"old", "new"}  # (only the new version is shown)
        if pending_break:
            if first.tag == qn("w:tbl"):
                # A table cannot start a page by itself: an empty paragraph carries the break.
                spacer = OxmlElement("w:p")
                first.addprevious(spacer)
                first = spacer
                ppr = spacer.get_or_add_pPr()
                spacing = OxmlElement("w:spacing")
                spacing.set(qn("w:before"), "0")
                spacing.set(qn("w:after"), "0")
                spacing.set(qn("w:line"), "20")
                spacing.set(qn("w:lineRule"), "exact")
                ppr.append(spacing)
            self.pending_break |= self.page_break_before(first, pending_break)
        if pending_notes:
            from docx.text.paragraph import Paragraph as DocxParagraph

            target = first if first.tag == qn("w:p") else first.find(".//" + qn("w:p"))
            at_start = True
            if target is not None and paragraph_style_id(target).startswith("Heading"):
                # The item was removed at the end of the previous section: anchor the note there.
                previous = [child for child in content_children(body)[:before] if child.tag in (qn("w:p"), qn("w:tbl"))]
                if previous and previous[-1].tag == qn("w:tbl"):
                    previous = previous[-1].findall(".//" + qn("w:p"))  # (into its last cell)
                if previous:
                    target, at_start = previous[-1], False
            if target is not None:
                for note in pending_notes:
                    # Where the removed item was: at the start of what follows it (or at the end
                    # of the preceding paragraph if a new section follows).
                    self.comment(DocxParagraph(target, self.doc._body), note, at_start=at_start)

    def page_break_before(self, p, versions: set) -> set:
        """Page break before paragraph P in the given versions (see below for a break of one
        version only). Returns the versions in which P does not exist: their break belongs to
        the next block."""
        ppr = p.get_or_add_pPr()
        rpr = ppr.find(qn("w:rPr"))
        mark = next((rev for rev, tag in REVISION_TAGS.items() if rpr is not None and rpr.find(qn(tag)) is not None),
                    None)
        present = paragraph_versions(p) if self.tracked else {"old", "new"}
        leftover = versions - present
        versions = versions & present
        if not versions:
            return leftover
        if present == {"old"}:
            # A paragraph of the old version only gets no break: LibreOffice passes the break of a
            # deleted paragraph on to the following one when accepting, and a tracked property
            # change makes it merge a following heading into it (the rejected document may lack
            # this page break; the accepted one is exact).
            return leftover
        if ppr.find(qn("w:pageBreakBefore")) is not None:
            return leftover  # (already breaks in the current version)
        change = ppr.find(qn("w:pPrChange"))
        if change is not None and "old" in versions:
            # Already a tracked property change (of the style, say): the old version's break
            # joins it.
            old_ppr = change.find(qn("w:pPr"))
            if old_ppr is not None and old_ppr.find(qn("w:pageBreakBefore")) is None:
                old_ppr.append(OxmlElement("w:pageBreakBefore"))
                reorder_children(old_ppr, PPR_ORDER)
        if "new" in versions or not self.tracked or mark in ADDED:
            set_page_break_before(p)
        # A break of only one version (an unchanged paragraph after an added or removed figure,
        # which has a page of its own) is not a tracked change: Word and LibreOffice would mark
        # the whole unchanged paragraph as reformatted. The new version's break is kept, the old
        # version's dropped (only the page layout of a rejected document differs).
        return leftover

    def write_block(self, block, nxt) -> None:
        if isinstance(block, Paragraph):
            self.write_paragraph(block)
        elif isinstance(block, Heading):
            self.write_heading(block)
        elif isinstance(block, Equation):
            self.write_equation(block)
        elif isinstance(block, Figure):
            self.write_figure(block, has_following=nxt is not None)
        elif isinstance(block, Table):
            self.write_table(block)
        elif isinstance(block, TableOfContents):
            self.write_toc(block)
        elif isinstance(block, ListBlock):
            self.write_list(block)

    def plain_revisions(self, block):
        """Tracked mode: a move written as a deletion and an insertion (see compare.detect_moves)."""
        if not self.tracked or not block.move or block.move_native:
            return block
        mapping = {"movefrom": "del", "moveto": "ins"}

        def retag(runs):
            out = []
            for run in runs or []:
                if hasattr(run, "rev") and run.rev in mapping:
                    run = replace(run, rev=mapping[run.rev])
                out.append(run)
            return out

        block = copy(block)
        block.runs = retag(block.runs)
        if isinstance(block, Heading) and block.number_runs is not None:
            block.number_runs = retag(block.number_runs)
        block.mark_rev = mapping.get(block.mark_rev, block.mark_rev)
        block.move = None
        return block

    def fit_bibliography_labels(self, blocks: list) -> None:
        """Make the hanging indent of labeled references as wide as the widest label (as
        thebibliography does), counting both versions of a tracked label such as [1][2]."""
        widest = 0.0
        for block in blocks:
            if not (isinstance(block, Paragraph) and block.style == "Bibliography Labeled"):
                continue
            label = ""
            for run in block.runs:
                if not isinstance(run, Text) or (not self.tracked and run.rev in REMOVED):
                    continue
                label += run.text
                if "\t" in label:
                    break
            if "\t" in label:
                widest = max(widest, text_width(label.split("\t", 1)[0], BASE_SIZE))
        if widest:
            fmt = self.doc.styles["Bibliography Labeled"].paragraph_format
            indent = Pt(max(30.0, widest + 6.0))
            fmt.left_indent = indent
            fmt.first_line_indent = -indent

    def write_paragraph(self, block: Paragraph):
        block = self.plain_revisions(block)
        style = block.style if block.style in [s.name for s in self.doc.styles] else "Normal"
        if block.list_kind:
            style = "List Paragraph"
        children = content_children(self.doc.element.body)
        after_table = bool(children) and children[-1].tag == qn("w:tbl") and not _is_layout_table(children[-1])
        paragraph = self.doc.add_paragraph(style=style)
        fmt = paragraph.paragraph_format
        if after_table:
            fmt.space_before = Pt(10)  # (a table with its caption above: separated as a float)
        if block.list_kind:
            self.list_format(paragraph, block)
        elif block.style in ("Normal",):
            paragraph.alignment = ALIGN.get(block.align, WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif block.style in ("Author", "Affiliation", "Title") and block.align in ALIGN:
            paragraph.alignment = ALIGN[block.align]
        elif block.style == "Caption" and block.align == "center":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER  # (a short caption of its own)
        if block.indent and not block.list_kind:
            fmt.first_line_indent = PARINDENT
        if block.hanging and block.style == "Normal" and not block.list_kind:
            fmt.left_indent = Pt(14.3)
            fmt.first_line_indent = Pt(-14.3)
        if block.quote and not block.list_kind:
            # quote/quotation/verse: indented on both sides by \leftmargin (2.5em) per level.
            fmt.left_indent = Pt(QUOTE_INDENT_PT * block.quote + (14.3 if block.hanging else 0))
            fmt.right_indent = Pt(QUOTE_INDENT_PT * block.quote)
        if block.space_before:
            fmt.space_before = Pt(block.space_before)
        size = SIZE_PT.get(block.size) if block.size != "normal" and block.style in ("Normal",) else None
        # (Runs in the document's language among runs set in another one - an English phrase
        # in a German sentence - keep it; a paragraph set in another language as a whole - a
        # bibliography entry - is in that language throughout.)
        self.paragraph_language = block.language if any(
            isinstance(r, Text) and r.lang == block.language for r in block.runs) else None
        try:
            self.add_runs(paragraph, block.runs, size)
        finally:
            self.paragraph_language = None
        if block.language in LANGUAGE_TAGS:
            # (Proofing language of a passage in another language: an English abstract, say.)
            for r in paragraph._p.iter(qn("w:r")):
                rpr = r.find(qn("w:rPr"))
                if rpr is None or rpr.find(qn("w:lang")) is None:
                    set_run_language(r, LANGUAGE_TAGS[block.language])
        if block.keep_with_next:
            fmt.keep_with_next = True  # (a panel heading with its table)
        if block.box:
            frame_paragraph(paragraph, block.box)
        if block.old_list is not None:
            self.list_change(paragraph, block)
        self.mark_paragraph(paragraph, block.mark_rev)
        self.mark_move(paragraph, block)
        if self.highlight and block.move_note and block.mark_rev == "moveto":
            self.comment(paragraph, block.move_note, at_start=True)
        elif self.tracked and block.move_note and block.mark_rev == "moveto" and "while moving" in block.move_note:
            # (Word shows the move; what it lost is only struck out at its old place: a comment
            # below it - not in the moved paragraph, whose move LibreOffice could then not reject.)
            what = "list item" if block.list_kind else "paragraph"
            self.comment(self.tiny_paragraph(), block.move_note.replace(f"This {what}", f"The {what} above", 1))
        for note in block.notes or []:
            # (A moved sentence: blue in highlight mode, a deletion plus an insertion when tracked.)
            self.comment(paragraph, note)
        if block.bookmark:
            self.add_bookmark(paragraph, block.bookmark)
        return paragraph

    def list_format(self, paragraph, block: Paragraph) -> None:
        level = min(block.list_level, 8)
        fmt = paragraph.paragraph_format
        if block.list_kind == "description":
            fmt.left_indent = Pt(18 * (level + 1))
            fmt.first_line_indent = Pt(-18) if block.list_item_start else Pt(0)
            paragraph.alignment = ALIGN.get(block.align, WD_ALIGN_PARAGRAPH.JUSTIFY)
            label = ""  # (of the new version: the tab stop stays when the changes are accepted)
            for run in block.runs if block.list_item_start else []:
                if not isinstance(run, Text) or "\t" in label:
                    break
                if run.rev not in REMOVED:
                    label += run.text
            if "\t" in label:
                # \item[label] of an itemize/enumerate: the tab goes to the text column, or
                # (a wider label) half an em after the label.
                width = text_width(label.split("\t", 1)[0], BASE_SIZE)
                if width + 5 > 18:
                    tabs = paragraph._p.get_or_add_pPr()
                    element = OxmlElement("w:tabs")
                    tab = OxmlElement("w:tab")
                    tab.set(qn("w:val"), "left")
                    tab.set(qn("w:pos"), str(int((18 * level + width + 5) * 20)))
                    element.append(tab)
                    tabs.append(element)
                    reorder_children(tabs, PPR_ORDER)
            return
        paragraph.alignment = ALIGN.get(block.align, WD_ALIGN_PARAGRAPH.JUSTIFY)
        if block.list_item_start:
            ppr = paragraph._p.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), str(level))
            num_id = OxmlElement("w:numId")
            num_id.set(qn("w:val"), str(self.list_num_id(block.list_kind, block.list_id, level, block.list_format)))
            num_pr.append(ilvl)
            num_pr.append(num_id)
            ppr.append(num_pr)
            reorder_children(ppr, PPR_ORDER)
        else:
            fmt.left_indent = Emu(int(LIST_INDENT_TWIPS * (level + 1) * 635))

    def write_heading(self, block: Heading) -> None:
        block = self.plain_revisions(block)
        level = min(max(block.level, 1), 5)
        paragraph = self.doc.add_paragraph(style=f"Heading {level}")
        runs = heading_runs(block, unbold=True)
        if self.highlight and block.old_level is not None:
            # (A section that became a subsection, say: its title is shown as reformatted.)
            runs = [replace(r, rev="fmt") if isinstance(r, Text) and r.rev is None and r.text.strip() else r
                    for r in runs]
        self.add_runs(paragraph, runs)
        self.mark_paragraph(paragraph, block.mark_rev)
        if block.old_level is not None and block.mark_rev is None:
            # (A section that became a subsection, say: rejecting restores the old level.)
            self.style_change(paragraph, f"Heading {min(max(block.old_level, 1), 5)}")
        self.mark_move(paragraph, block)
        if self.highlight and block.move_note and block.mark_rev == "moveto":
            self.comment(paragraph, block.move_note, at_start=True)
        if block.bookmark:
            self.add_bookmark(paragraph, block.bookmark)

    def write_toc(self, block: TableOfContents) -> None:
        toc_rev = block.rev if block.rev in ("ins", "del") else None
        toc_levels = getattr(block, "levels", 3) or 3
        if self.highlight and toc_rev == "del":
            return
        title = self.doc.add_paragraph(style="TOC Heading")
        self.add_runs(title, [Text(block.title, Style(), toc_rev)])
        self.mark_paragraph(title, toc_rev)
        entries = block.entries if block.entries is not None else self.toc_entries
        if self.highlight and block.entries is not None:
            entries = accept_toc_deletions(entries)
        if not entries:
            return
        paragraphs = []
        for entry in entries:
            paragraph = self.doc.add_paragraph(style=f"toc {min(max(entry.level, 1), 3)}")
            if entry.old_level is not None and entry.mark_rev is None:
                self.style_change(paragraph, f"toc {min(max(entry.old_level, 1), 3)}")
            link = OxmlElement("w:hyperlink")
            if entry.bookmark:
                link.set(qn("w:anchor"), entry.bookmark)
            link.set(qn("w:history"), "1")
            paragraph._p.append(link)
            # (A table of contents that only exists in one version: its entries are tracked too.)
            self.mark_paragraph(paragraph, entry.mark_rev or toc_rev)
            for run in entry.runs:
                if not isinstance(run, Text) or not run.text:
                    continue
                rev = run.rev or toc_rev
                if self.highlight and rev == "del":
                    continue
                r = self.make_run(run.text, Style(italic=run.style.italic), None)
                if self.highlight and rev in ("ins", "fmt"):
                    self.apply_highlight(r)
                self.place(link, r, rev if rev in ("ins", "del") else None)
            for wrapper in link.iter(qn("w:ins"), qn("w:del")):
                wrapper.attrib.pop("_merge", None)
            paragraphs.append(paragraph)
        # Wrap the entries in a TOC field so that Word can regenerate it (with page numbers).
        first, last = paragraphs[0], paragraphs[-1]
        begin_runs = []
        for kind, content in (("begin", None), ("instr", f'TOC \\o "1-{toc_levels}" \\h \\z \\u'), ("separate", None)):
            r = OxmlElement("w:r")
            if kind == "instr":
                instr = OxmlElement("w:instrText")
                instr.set(qn("xml:space"), "preserve")
                instr.text = f" {content} "
                r.append(instr)
            else:
                char = OxmlElement("w:fldChar")
                char.set(qn("w:fldCharType"), kind)
                r.append(char)
            begin_runs.append(r)
        end = OxmlElement("w:r")
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), "end")
        end.append(char)
        if self.tracked and toc_rev:
            # A table of contents that only exists in one version: its field codes too.
            wrapped = []
            for r in begin_runs:
                wrapper = self.revision_element("w:ins" if toc_rev == "ins" else "w:del")
                if toc_rev == "del":
                    to_deleted(r)
                wrapper.append(r)
                wrapped.append(wrapper)
            begin_runs = wrapped
            wrapper = self.revision_element("w:ins" if toc_rev == "ins" else "w:del")
            wrapper.append(end)
            end = wrapper
        ppr = first._p.find(qn("w:pPr"))
        index = 1 if ppr is not None else 0
        for offset, r in enumerate(begin_runs):
            first._p.insert(index + offset, r)
        last._p.append(end)

    def equation_rev(self, rev: str | None) -> str | None:
        """Tracked mode shows a moved equation as deleted and inserted (no move markup for the
        rows of equation tables)."""
        if self.tracked and rev in ("movefrom", "moveto"):
            return "del" if rev == "movefrom" else "ins"
        return rev

    def write_equation(self, block: Equation) -> None:
        self.separate_tables()
        rev = self.equation_rev(block.rev)
        marked = self.highlight and rev in ("ins", "fmt", "moveto")
        frame = self.changes.move_frame_color if rev == "moveto" else (self.changes.frame_color if self.changes else None)
        if block.number is None:
            paragraph = self.doc.add_paragraph(style="Equation")
            self.place_math(paragraph._p, block.omml, rev, display=True)
            self.mark_paragraph(paragraph, rev)
            if marked:
                box_paragraph(paragraph, frame)
            return
        # Numbered equations: a borderless three-column table keeps the formula in display
        # mode (centered) with the number flush right, as in LaTeX.
        table = self.doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = (Inches(0.5), Inches(TEXT_WIDTH_IN - 1.0), Inches(0.5))
        set_table_borders(table, none=True)
        set_table_width(table, sum(w for w in widths))
        set_table_layout_fixed(table)
        set_grid_widths(table, [w.inches for w in widths])
        for cell, width in zip(table.rows[0].cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 0, 0)
        middle = table.rows[0].cells[1].paragraphs[0]
        middle.style = self.doc.styles["Equation"]
        self.place_math(middle._p, block.omml, rev, display=True)
        number = table.rows[0].cells[2].paragraphs[0]
        number.style = self.doc.styles["Equation"]
        number.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if block.number_runs is not None:
            self.add_runs(number, block.number_runs)
        elif block.number:
            self.add_runs(number, [Text(equation_label(block), Style(), rev)])
        left = table.rows[0].cells[0].paragraphs[0]
        left.style = self.doc.styles["Equation"]
        self.mark_equation_row(table.rows[0], rev)
        if rev in ("ins", "del"):
            for paragraph in (left, middle, number):
                self.mark_paragraph(paragraph, rev)
        if marked:
            for edge in ("top", "left", "bottom", "right"):
                set_cell_border(table.rows[0].cells[1], edge, 12, frame)
        mark_table_as_layout(table)

    def mark_equation_row(self, row, rev: str | None) -> None:
        # The row revision makes Word remove the row when the change is accepted (deleted
        # equation) or rejected (inserted equation), so that a paragraph split by a new equation
        # rejoins. (LibreOffice paints such rows light cyan and draws formulas in the automatic
        # color, which is hard to read in its dark mode; the highlight mode avoids this.)
        self.mark_row(row, rev)

    def write_equation_group(self, rows: list) -> None:
        self.separate_tables()
        """Rows of an align environment: one table, aligned at the & columns, numbered per row."""
        nparts = max(len(r.parts) for r in rows)
        aligns = (rows[0].aligns or "rl")
        pattern = [aligns[i % len(aligns)] if len(aligns) < 3 else aligns[min(i, len(aligns) - 1)]
                   for i in range(nparts)]
        pad, num = 0.35, 0.6
        available = TEXT_WIDTH_IN - pad - num
        part_widths = [available / nparts] * nparts
        table = self.doc.add_table(rows=len(rows), cols=nparts + 2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [pad] + part_widths + [num]
        set_table_borders(table, none=True)
        set_table_width(table, float(sum(widths)))
        set_table_layout_fixed(table)
        set_grid_widths(table, widths)
        jc = {"r": "right", "l": "left", "c": "center"}
        for row_index, eq in enumerate(rows):
            tr = table.rows[row_index]
            set_cant_split(tr)
            cells = tr.cells
            for col, width in enumerate(widths):
                cells[col].width = Inches(width)
                cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cells[col], 0, 0)
                paragraph = cells[col].paragraphs[0]
                paragraph.style = self.doc.styles["Equation"]
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
            rev = self.equation_rev(eq.rev)
            for part_index, part in enumerate(eq.parts):
                paragraph = cells[part_index + 1].paragraphs[0]
                align = eq.row_align or pattern[part_index]
                paragraph.alignment = ALIGN.get(jc[align], WD_ALIGN_PARAGRAPH.CENTER)
                if part:
                    self.place_math(paragraph._p, part, rev, display=True, justify=jc[align])
            number = cells[-1].paragraphs[0]
            number.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if eq.number_runs is not None:
                self.add_runs(number, eq.number_runs)
            elif eq.number:
                self.add_runs(number, [Text(equation_label(eq), Style(), rev)])
            elif self.tracked and rev in ("ins", "del"):
                # An invisible tracked character: LibreOffice ignores tracked formulas and would
                # show an inserted/deleted unnumbered row (\notag) without any change bar.
                self.add_runs(number, [Text("\u200b", Style(), rev)])
            self.mark_equation_row(tr, rev)
            if rev in ("ins", "del"):
                for cell in cells:
                    self.mark_paragraph(cell.paragraphs[0], rev)
        if self.highlight:
            # A frame around the changed (or moved) rows (formulas ignore the shading in LibreOffice).
            colors = [self.changes.frame_color if eq.rev in ("ins", "fmt") else
                      self.changes.move_frame_color if eq.rev == "moveto" else None for eq in rows]
            for row_index, color in enumerate(colors):
                if color is None:
                    continue
                cells = table.rows[row_index].cells[1:-1]
                for cell in cells:
                    if row_index == 0 or colors[row_index - 1] != color:
                        set_cell_border(cell, "top", 12, color)
                    if row_index == len(rows) - 1 or colors[row_index + 1] != color:
                        set_cell_border(cell, "bottom", 12, color)
                set_cell_border(cells[0], "left", 12, color)
                set_cell_border(cells[-1], "right", 12, color)
        mark_table_as_layout(table)

    def space_after_last(self, points: float) -> None:
        """Separate a float from the preceding text paragraph, as LaTeX does."""
        body = self.doc.element.body
        children = [c for c in body if c.tag != qn("w:sectPr")]
        if children and children[-1].tag == qn("w:p"):
            paragraph = self.doc.paragraphs[-1]
            spacing = children[-1].find(f"{qn('w:pPr')}/{qn('w:spacing')}")
            if spacing is not None and spacing.get(qn("w:lineRule")) == "exact" and int(spacing.get(qn("w:line"), "999")) <= 40:
                return  # (a minimal separator between two tables stays minimal)
            if paragraph._p is children[-1] and paragraph.style.name in ("Normal", "Equation"):
                paragraph.paragraph_format.space_after = Pt(points)

    def picture_paragraph(self, image: Path, width_in: float, height_in: float, caption_runs: list,
                          descr: str, rev: str | None, frame: bool = False, frame_color: str | None = None):
        paragraph = self.doc.add_paragraph(style="Figure")
        r = self.picture_run(paragraph, image, width_in, height_in, caption_runs, descr, frame, frame_color)
        self.place(paragraph._p, r, rev)
        for wrapper in paragraph._p.iter(qn("w:ins"), qn("w:del")):
            wrapper.attrib.pop("_merge", None)
        self.mark_paragraph(paragraph, rev)
        if getattr(self, "figure_break_done", False):
            self.figure_break_done = False  # (the break is before the text above the figure)
        else:
            self.page_break_before(paragraph._p, versions_of(rev))
        return paragraph

    def picture_run(self, paragraph, image: Path, width_in: float, height_in: float, caption_runs: list,
                    descr: str, frame: bool = False, frame_color: str | None = None):
        """A detached w:r holding the picture, scaled to fit the page together with its caption."""
        caption_chars = sum(len(r.text) for r in caption_runs if isinstance(r, Text)) + 12
        caption_lines = math.ceil(caption_chars / 100) if caption_runs else 0
        caption_height = caption_lines * 0.19 + 0.45
        max_height = TEXT_HEIGHT_IN - caption_height - 0.3
        scale = min(1.0, TEXT_WIDTH_IN / width_in, max_height / height_in)
        run = paragraph.add_run()
        run.add_picture(str(image), width=Emu(int(Inches(width_in * scale))), height=Emu(int(Inches(height_in * scale))))
        doc_pr = run._r.find(".//" + qn("wp:docPr"))
        if doc_pr is not None:
            doc_pr.set("descr", descr)
        if frame:
            for sppr in run._r.iter(qn("pic:spPr")):
                add_shape_outline(sppr, frame_color or self.changes.frame_color)
        paragraph._p.remove(run._r)
        return run._r

    def changed_picture_paragraph(self, old: Figure, new: Figure, caption_runs: list, descr: str):
        """Tracked mode, changed figure: one paragraph with the previous picture (deleted) and the
        new picture (inserted) - a wide picture wraps onto a line of its own, and no separator is
        left behind by accepting or rejecting. The paragraph itself and its page break exist in
        both versions, so accepting or rejecting keeps the layout."""
        paragraph = self.doc.add_paragraph(style="Figure")
        old_run = self.picture_run(paragraph, old.image, old.width_in or TEXT_WIDTH_IN, old.height_in or TEXT_WIDTH_IN,
                                   caption_runs, descr + " (previous version)")
        self.place(paragraph._p, old_run, "del")
        new_run = self.picture_run(paragraph, new.image, new.width_in or TEXT_WIDTH_IN, new.height_in or TEXT_WIDTH_IN,
                                   caption_runs, descr)
        self.place(paragraph._p, new_run, "ins")
        # An invisible unchanged character ends the paragraph: when a tracked picture is the last
        # thing in a paragraph, LibreOffice joins the paragraph with the caption after it on
        # accepting or rejecting.
        paragraph._p.append(self.make_run("\u200b", Style(), None))
        for wrapper in paragraph._p.iter(qn("w:ins"), qn("w:del")):
            wrapper.attrib.pop("_merge", None)
        if getattr(self, "figure_break_done", False):
            self.figure_break_done = False  # (the break is before the text above the figure)
        else:
            set_page_break_before(paragraph._p)
        return paragraph

    def write_figure(self, block: Figure, has_following: bool) -> None:
        rev = block.rev
        caption_runs = block.caption_runs if block.caption_runs is not None else (
            [Text(figure_label(block))] + list(block.caption) if block.caption else [])
        noun = block.name or block.kind.capitalize()
        descr = f"{noun} {block.number}" if block.number else noun
        changed = block.change == "changed" and block.old is not None
        rendered = block.image is not None and block.image.is_file()
        old_rendered = changed and block.old.image is not None and block.old.image.is_file()
        old_shown = changed and self.tracked and rendered and old_rendered
        # (Only one of the versions could be rendered: both shown, the other as a placeholder.)
        half = changed and self.tracked and rendered != old_rendered
        frame = self.highlight and (changed or rev == "ins")
        moved = rev == "ins" and bool(block.move_note)
        if old_shown:
            paragraph = self.changed_picture_paragraph(block.old, block, caption_runs, descr)
        elif half:
            paragraph = self.doc.add_paragraph(style="Figure")
            for version, version_rev in ((block.old, "del"), (block, "ins")):
                if version.image is not None and version.image.is_file():
                    r = self.picture_run(paragraph, version.image, version.width_in or TEXT_WIDTH_IN,
                                         version.height_in or TEXT_WIDTH_IN, caption_runs, descr)
                else:
                    r = self.make_run(f"[Figure not rendered: {version.error or 'unknown error'}]",
                                      Style(italic=True), None)
                self.place(paragraph._p, r, version_rev)
            paragraph._p.append(self.make_run("\u200b", Style(), None))  # (see changed_picture_paragraph)
            for wrapper in paragraph._p.iter(qn("w:ins"), qn("w:del")):
                wrapper.attrib.pop("_merge", None)
            if getattr(self, "figure_break_done", False):
                self.figure_break_done = False
            else:
                set_page_break_before(paragraph._p)
        elif rendered:
            paragraph = self.picture_paragraph(block.image, block.width_in or TEXT_WIDTH_IN,
                                               block.height_in or TEXT_WIDTH_IN, caption_runs, descr, rev,
                                               frame=frame, frame_color=self.changes.move_frame_color
                                               if moved and self.changes else None)
        else:
            paragraph = self.doc.add_paragraph(style="Figure")
            message = f"[Figure not rendered: {block.error or 'unknown error'}]"
            self.add_runs(paragraph, [Text(message, Style(italic=True), rev)])
            self.mark_paragraph(paragraph, rev)
            if getattr(self, "figure_break_done", False):
                self.figure_break_done = False
            else:
                self.page_break_before(paragraph._p, versions_of(rev))
        caption = None
        if caption_runs:
            caption = self.doc.add_paragraph(style="Caption")
            self.add_runs(caption, caption_runs)
            center_short_caption(caption, caption_runs, self.tracked)
            self.mark_paragraph(caption, rev)
        else:
            paragraph.paragraph_format.keep_with_next = False
        if changed:
            what = f"The size of this {block.kind} was changed" if block.resized else f"This {block.kind} was changed"
            if old_shown:
                note = f"{what}. The previous version is shown as deleted before it."
            elif half:
                note = (f"{what}. LaTeX could not render one of the two versions (shown as a placeholder); "
                        "the previous version is shown as deleted before the new one.")
            elif rendered or not self.tracked:
                note = f"{what}."
            else:
                note = f"This {block.kind} was changed (figures were not rendered, so neither version can be shown)."
            # In tracked mode the comment must not sit in a tracked-inserted paragraph (LibreOffice
            # then fails to reject the insertion): it goes on the caption, which exists in both versions.
            if not self.tracked:
                self.comment(paragraph, note)
            elif caption is not None:
                self.comment(caption, note, at_start=True)
            elif not (old_shown or half):
                self.comment(paragraph, note)
        elif rev == "ins" and not self.tracked:
            # (In tracked mode a new figure is visibly a tracked insertion.)
            self.comment(paragraph, block.move_note if moved else f"This {block.kind} is new.")
        elif moved and self.tracked:
            # (On an untracked 1-pt paragraph right below the moved figure: see write_table.)
            self.comment(self.tiny_paragraph(), block.move_note.replace(f"This {block.kind}", f"The {block.kind} above", 1))
        # Each figure gets a page of its own: "page break before" on its first paragraph and on
        # the content that follows (unlike a paragraph holding a manual page break, this never
        # produces an empty page or a stretched last line).
        self.pending_break = versions_of(block.rev) if has_following else set()

    def comment(self, paragraph, text: str, at_start: bool = False, anchor=None) -> None:
        if self.changes is None or not self.changes.comments:
            return
        # Anchor on a run of its own: the picture run may sit inside a w:ins/w:del wrapper,
        # which python-docx's comment API cannot address.
        if anchor is not None:
            from docx.text.run import Run

            run = Run(anchor, paragraph)
        else:
            run = paragraph.add_run("")
        if at_start and anchor is None:
            # (After the comments already anchored at the start: several keep their order.)
            p = paragraph._p
            p.remove(run._r)
            children = list(p)
            position = 1 if children and children[0].tag == qn("w:pPr") else 0
            while position < len(children) and _comment_markup(children[position]):
                position += 1
            p.insert(position, run._r)
        author = self.changes.author
        try:
            comment = self.doc.add_comment([run], text=text, author=author,
                                           initials="".join(w[0] for w in author.split()[:3]).upper() or "D")
            # The same (Word-style local) time stamp as the tracked changes.
            comment._comment_elm.set(qn("w:date"), self.changes.date_text)
            for comment_paragraph in comment.paragraphs:
                for comment_run in comment_paragraph.runs:
                    comment_run.font.color.rgb = COMMENT_TEXT_COLOR
                    # (English explanations quoting text in the document's language: no spelling marks)
                    comment_run.font.no_proof = True
        except Exception:  # pragma: no cover - comments are a nicety
            pass

    def separate_tables(self, gap: float = 0) -> None:
        """Two tables in a row (a data table, an equation) would merge into one in Word and
        LibreOffice: an untracked empty paragraph between them (GAP points high below)."""
        children = content_children(self.doc.element.body)
        if children and children[-1].tag == qn("w:tbl"):
            self.tiny_paragraph().paragraph_format.space_after = Pt(gap)

    def write_table(self, block: Table) -> None:
        self.separate_tables(gap=8)  # (two tables, not one)
        self.space_after_last(10)
        rev = block.rev
        caption_runs = block.caption_runs if block.caption_runs is not None else (
            [Text(table_label(block))] + list(block.caption) if block.caption else [])
        # A table whose columns changed is shown as the deleted old table followed by the
        # inserted new one, with one caption shared by both versions.
        replacement = self.tracked and (block.replaced or (rev == "ins" and block.highlight_rows is not None))
        caption_rev = None if replacement else rev
        caption = None
        caption_moved = block.old_caption_runs is not None and rev is None and not replacement
        if caption_moved:
            if self.tracked:
                # (The old caption deleted on its side of the table, the new one inserted on the other.)
                caption_runs = runs_with_rev([Text(table_label(block))] + list(block.caption), "ins")
                if not block.caption_above:
                    self.write_caption(block.old_caption_runs, "del", keep_with_next=True)
                caption_rev = "ins"
            else:
                caption_runs = [replace(r, rev="ins", moved=True) if isinstance(r, Text) and r.rev is None else r
                                for r in caption_runs]
        if caption_runs and block.caption_above:
            caption = self.write_caption(caption_runs, caption_rev, keep_with_next=True)
        intro = block.intro_runs if block.intro_runs is not None else block.intro
        if intro:
            self.write_float_text(intro, rev, block.text_align, keep_with_next=True)
        rows = block.rows
        row_revs = block.row_revs or [rev] * len(rows)
        in_place = self.tracked and block.tracked_rows is not None  # (rows reordered within their groups)
        if in_place:
            rows, row_revs = block.tracked_rows, [rev] * len(block.tracked_rows)
        if self.highlight:
            kept = [(row, r) for row, r in zip(rows, row_revs) if r != "del"]
            rows = [row for row, _ in kept]
            row_revs = [r for _, r in kept]
        if not rows:
            return
        ncols = block.ncols or max(sum(c.colspan for c in r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, none=True)
        size = SIZE_PT.get(block.size) if block.size != "normal" else None
        widths = estimate_column_widths(rows, ncols, size or BASE_SIZE)
        set_table_width(table, sum(widths))
        set_grid_widths(table, widths)
        header = block.header_rows if not any(row_revs[:block.header_rows]) else 0
        for row_index, row in enumerate(rows):
            tr = table.rows[row_index]
            set_cant_split(tr)
            if row_index < header:
                tr._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            col = 0
            for cell in row:
                if col >= ncols:
                    break
                target = tr.cells[col]
                if cell.colspan > 1:
                    target = target.merge(tr.cells[min(col + cell.colspan, ncols) - 1])
                paragraph = target.paragraphs[0]
                paragraph.alignment = ALIGN.get(cell.align, WD_ALIGN_PARAGRAPH.LEFT)
                paragraph.paragraph_format.first_line_indent = Pt(0)
                if not cell.placeholder:
                    runs = cell.runs
                    if self.highlight and cell.same_label:  # (the group's label, unchanged)
                        runs = [replace(r, rev=None) if getattr(r, "rev", None) == "ins" else r for r in runs]
                    self.add_runs(paragraph, runs, size)
                if cell.fill:
                    fill_cell(target, cell.fill)
                if row_revs[row_index] in ("ins", "del"):
                    self.mark_paragraph(paragraph, row_revs[row_index])
                if cell.rowspan > 1 or (cell.merge_start and row_revs[row_index] == "del"):
                    # Centered like \multirow, except across tracked rows: LibreOffice paints
                    # those light cyan, where a light (dark mode) label would be half hidden.
                    span = row_revs[row_index:row_index + cell.rowspan]
                    alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if self.tracked and any(r in ("ins", "del") for r in span):
                        # The label goes into an unchanged row of the merged cell.
                        if span[0] not in ("ins", "del"):
                            alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                        elif span[-1] not in ("ins", "del"):
                            alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
                    target.vertical_alignment = alignment
                    set_vmerge(target, "restart")
                elif cell.placeholder:
                    set_vmerge(target, None)
                col += cell.colspan
            # Each cell is as wide as the grid columns it spans (merged \multicolumn cells too).
            for cell_ref, first, span in row_cells(table.rows[row_index]):
                tc_pr = cell_ref._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.insert(0, tc_w)
                tc_w.set(qn("w:w"), str(int(round(sum(widths[first:first + span]) * 1440))))
                tc_w.set(qn("w:type"), "dxa")
                reorder_tcpr(tc_pr)
            self.mark_row(tr, row_revs[row_index])
        # A short table (as a LaTeX tabular) is not split across pages: every row keeps with the
        # next one (and the last row with a caption below the table).
        if len(rows) <= 25:
            keep_last = bool(caption_runs) and not block.caption_above
            for row_index, row in enumerate(table.rows):
                if row_index < len(rows) - 1 or keep_last:
                    for cell_paragraph in row._tr.iter(qn("w:p")):
                        cell_paragraph.get_or_add_pPr().keepNext_val = True
        # LaTeX rules: \hline above row i  ->  top border of row i (bottom border of the last row).
        rules = block.rules
        if in_place:
            rules = block.tracked_rules[0]
        elif self.highlight and block.row_revs:
            rules = remap_rules(block.rules, block.row_revs)
        elif self.tracked:
            rules = rules | block.old_rules
        # (Borders go on the row's own cells: python-docx reports the first cell of a vertical
        # merge for every row it spans, which would put a rule below a merge on the wrong cell.)
        for rule in rules:
            if rule < len(rows):
                for cell, _, _ in row_cells(table.rows[rule]):
                    set_cell_border(cell, "top")
            elif rows:
                for cell, _, _ in row_cells(table.rows[len(rows) - 1]):
                    set_cell_border(cell, "bottom")
                kept = [i for i, rev in enumerate(row_revs) if rev != "del"]
                if self.tracked and row_revs[-1] == "del" and kept:
                    # (The last row was deleted: the rule below the table stays after accepting.)
                    for cell, _, _ in row_cells(table.rows[kept[-1]]):
                        set_cell_border(cell, "bottom")
        partial = block.tracked_rules[1] if in_place else (block.partial_rules or {})
        if self.highlight and block.row_revs:
            mapping = remap_rule_index(block.row_revs)
            partial = {mapping.get(row, row): spans for row, spans in partial.items()}
        for rule, spans in partial.items():
            target_row = min(rule, len(rows) - 1)
            edge = "top" if rule < len(rows) else "bottom"
            for cell, first, span in row_cells(table.rows[target_row]):
                if any(first <= b and first + span - 1 >= a for a, b in spans):
                    set_cell_border(cell, edge)
        for vcol in block.vrules:
            for row in table.rows:
                for cell, first, span in row_cells(row):
                    if first == vcol:
                        set_cell_border(cell, "left")
                    elif vcol >= ncols and first + span == ncols:
                        set_cell_border(cell, "right")
        for row in table.rows:
            for cell, _, _ in row_cells(row):
                set_cell_margins(cell, 40, 40)
        outro = block.outro_runs if block.outro_runs is not None else block.outro
        if outro:
            if len(rows) <= 25:  # (with the rows, as LaTeX keeps a tabular together)
                for cell_paragraph in table.rows[-1]._tr.iter(qn("w:p")):
                    cell_paragraph.get_or_add_pPr().keepNext_val = True
            self.write_float_text(outro, rev, block.text_align, keep_with_next=bool(caption_runs)
                                  and not block.caption_above)
        if caption_moved and self.tracked and block.caption_above:
            self.write_caption(block.old_caption_runs, "del")
        if caption_runs and not block.caption_above:
            caption = self.write_caption(caption_runs, caption_rev)
        elif replacement and block.replaced:
            # Separates the two tables (adjacent tables would merge) and stays in both versions.
            spacer = self.doc.add_paragraph()
            fmt = spacer.paragraph_format
            fmt.space_before = fmt.space_after = Pt(0)
            fmt.line_spacing = Pt(1)
            fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            if caption is None:
                caption = spacer
        elif not caption_runs and not (replacement and block.caption_above and block.highlight_caption):
            spacer = self.doc.add_paragraph()
            self.mark_paragraph(spacer, rev)
        moved_rows = sum(1 for row, row_rev in zip(rows, row_revs) if row_rev == "ins" and any(
            getattr(r, "moved", False) for c in row for r in c.runs))
        if self.tracked and moved_rows and not in_place and caption is not None and caption_rev is None:
            self.comment(caption, f"{moved_rows} row{'s' if moved_rows > 1 else ''} of this table "
                                  f"{'were' if moved_rows > 1 else 'was'} moved: shown as deleted at the old "
                                  "place and inserted at the new one.", at_start=True)
        if in_place and caption is not None and caption_rev is None:
            self.comment(caption, "Rows of this table were put in a different order within their groups; the "
                                  "changed cells show the old and the new order.", at_start=True)
        if caption_moved and caption is not None and not self.tracked:
            # (Tracked mode shows the old caption deleted on its side; a comment must not sit in the
            # inserted caption there.)
            self.comment(caption, "The caption of this table was moved from " +
                         ("below" if block.caption_above else "above") + " the table to " +
                         ("above" if block.caption_above else "below") + " it.", at_start=True)
        # (The note goes on the shared caption, or on the separator if the table has no caption.)
        noted = (block.caption_above or not block.highlight_caption) if block.replaced else not block.caption_above
        if self.tracked and rev == "ins" and block.move_note:
            # (Not in the tracked-inserted table or caption - LibreOffice then fails to reject the
            # insertion -: on an untracked 1-pt paragraph right below it.)
            self.comment(self.tiny_paragraph(), block.move_note.replace("This table", "The table above", 1))
        if self.highlight and rev == "ins" and block.move_note:
            target = caption if caption is not None else None
            if target is None:
                first_cell = table._tbl.find(".//" + qn("w:p"))
                from docx.text.paragraph import Paragraph as DocxParagraph

                target = DocxParagraph(first_cell, self.doc._body) if first_cell is not None else None
            if target is not None:
                self.comment(target, block.move_note, at_start=True)
        if replacement and caption is not None and noted:
            self.comment(caption, "Columns were added to or removed from this table, which cannot be shown as "
                                  "tracked changes: the previous version of the table is shown as deleted, "
                                  "followed by the new version.")

    def tiny_paragraph(self):
        """An untracked empty paragraph of minimal height (between two tables, under a moved float)."""
        spacer = self.doc.add_paragraph()
        fmt = spacer.paragraph_format
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing = Pt(1)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        return spacer

    def write_float_text(self, runs: list, rev: str | None, align: str, keep_with_next: bool = False):
        """Text that goes with a table, between its caption and the tabular ("Panel A", a note)."""
        if self.highlight and not any(isinstance(r, Text) and r.text.strip() and r.rev not in REMOVED for r in runs):
            return None
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.CENTER)
        fmt = paragraph.paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.space_before = Pt(2)
        fmt.space_after = Pt(4)
        fmt.keep_with_next = keep_with_next
        self.add_runs(paragraph, runs, SIZE_PT.get("small"))
        texts = {getattr(r, "rev", None) for r in runs if isinstance(r, Text) and r.text.strip()}
        mark = rev or (texts.pop() if len(texts) == 1 and texts <= {"ins", "del"} else None)
        self.mark_paragraph(paragraph, mark)
        return paragraph

    def write_caption(self, runs: list, rev: str | None, keep_with_next: bool = False):
        caption = self.doc.add_paragraph(style="Caption")
        self.add_runs(caption, runs)
        center_short_caption(caption, runs, self.tracked)
        self.mark_paragraph(caption, rev)
        if keep_with_next:
            caption.paragraph_format.keep_with_next = True
        return caption

    def write_list(self, block: ListBlock, level: int = 0) -> None:
        """Legacy list blocks (the converter now emits list paragraphs)."""
        for item in block.items:
            first = True
            for element in item:
                if isinstance(element, ListBlock):
                    self.write_list(element, level + 1)
                    continue
                element.list_kind = element.list_kind or ("number" if block.ordered else "bullet")
                element.list_level = level
                element.list_item_start = first
                self.write_paragraph(element)
                first = False

    def save(self, path: Path) -> None:
        self.doc.save(str(path))


# ---------------------------------------------------------------------- revision helpers


def quoted_text(text: str, limit: int = 25) -> str:
    """The start of a removed passage, for a comment."""
    words = text.split()
    return " ".join(words[:limit]) + (" …" if len(words) > limit else "")


def mark_deletion_sites(runs: list, quote: bool = False) -> list:
    """Highlight mode shows no deleted text: a short deletion that is not part of a replacement
    ("−0.24" -> "0.24", "< 0.001" -> "0.001") marks the word after it (before it, at the end)
    as changed, so that the edit is not invisible. A longer one (a phrase, a sentence) gets a
    comment quoting it, with QUOTE (in body text), or is not shown."""
    runs = list(runs)
    index = 0
    while index < len(runs):
        if getattr(runs[index], "rev", None) != "del" or getattr(runs[index], "moved", False):
            index += 1  # (a passage that moved elsewhere: the comment at its new place says so)
            continue
        end = index  # (consecutive deleted text, formulas and footnotes are one deletion)
        while end < len(runs) and getattr(runs[end], "rev", None) == "del" and not getattr(runs[end], "moved", False):
            end += 1
        deleted = "".join(r.text if isinstance(r, Text) else " x " for r in runs[index:end])
        neighbours = runs[max(0, index - 1):index] + runs[end:end + 1]
        deleted_words = len(re.findall(r"\w+", deleted))
        replaced = any(getattr(r, "rev", None) in ADDED for r in neighbours)
        if replaced:
            # (The insertion next to it shows the change - unless far more was removed than added:
            # a removed sentence next to a reworded word still gets its comment.)
            inserted = ""
            for step, k in ((1, end), (-1, index - 1)):  # (the insertion after it, and before it)
                while 0 <= k < len(runs) and getattr(runs[k], "rev", None) in ADDED:
                    inserted += runs[k].text if isinstance(runs[k], Text) else " x "
                    k += step
            if not (deleted_words > 5 and len(re.findall(r"\w+", inserted)) * 2 < deleted_words):
                index = end
                continue
        if not deleted.strip():
            index = end
            continue
        if len(deleted.strip()) > 40 or deleted_words > 5:
            kept = " ".join("".join(r.text for r in runs if isinstance(r, Text) and r.rev not in REMOVED).split())
            if " ".join(deleted.split()) in kept:
                index = end  # (the same text is still there - it moved within the paragraph)
                continue
            if quote and any(getattr(r, "rev", None) is None and isinstance(r, Text) and r.text.strip()
                             for r in runs):
                text = whole_words(runs, index, end, quoted_runs(runs[index:end]))
                runs.insert(index, CommentMark(f"Text was removed here (“{quoted_text(text)}”)."))
                end += 1
            index = end
            continue
        after = next((k for k in range(end, len(runs)) if isinstance(runs[k], Text) and runs[k].text.strip()
                      and runs[k].rev is None), None)
        before = next((k for k in range(index - 1, -1, -1) if isinstance(runs[k], Text) and runs[k].text.strip()
                       and runs[k].rev is None), None)
        if after is not None and before is not None and not any(c.isalnum() for c in runs[after].text.split()[0]):
            after = None  # (not a lone full stop: the word before the deletion)
        if after is not None and (before is None or after - end <= index - before):
            target = runs[after]
            match = re.match(r"(\s*\S+)(.*)", target.text, re.S)
            parts = [Text(match.group(1), target.style, "ins"), Text(match.group(2), target.style)]
            runs[after:after + 1] = [part for part in parts if part.text]
        elif before is not None:
            target = runs[before]
            match = re.match(r"(.*?)(\S+\s*)$", target.text, re.S)
            parts = [Text(match.group(1), target.style), Text(match.group(2), target.style, "ins")]
            runs[before:before + 1] = [part for part in parts if part.text]
            end += 1
        index = end + 1
    return runs


def accept_deletions(blocks: list) -> list:
    """Highlight mode: remove deleted content so that only the new text remains.

    Deleted runs disappear, paragraphs whose paragraph mark was deleted are merged with the
    following paragraph (as accepting the deletion in Word does: the merged paragraph keeps
    the properties of the surviving paragraph mark), and deleted figures, tables, equations
    and page breaks are dropped.
    """
    def keep(runs):
        return [r for r in runs if getattr(r, "rev", None) not in REMOVED]

    def text_of(runs) -> str:
        return "".join(r.text for r in runs or [] if isinstance(r, Text) and getattr(r, "rev", None) not in ADDED)

    # Removed equations: in each stretch of changed equations, those deleted beyond the number
    # inserted (a row of a display that lost a row while another row changed, too) get a note.
    noted: set[int] = set()
    position = 0
    while position < len(blocks):
        if not (isinstance(blocks[position], Equation) and blocks[position].rev in ("del", "ins")):
            position += 1
            continue
        end = position
        while end < len(blocks) and isinstance(blocks[end], Equation) and blocks[end].rev in ("del", "ins", None):
            end += 1
        stretch = range(position, end)
        deleted = [i for i in stretch if blocks[i].rev == "del"]
        inserted = [i for i in stretch if blocks[i].rev == "ins"]
        numbers = {blocks[i].number for i in inserted}
        surplus = len(deleted) - len(inserted)
        if surplus > 0:
            # (Prefer the equations whose number did not come back.)
            ranked = sorted(deleted, key=lambda i: (blocks[i].number in numbers, -i))
            noted.update(ranked[:surplus])
        position = end

    def footnote_text(note) -> str:
        return " ".join("".join(r.text for r in note.runs if isinstance(r, Text)).split())

    # Footnotes that reappear elsewhere (moved with their sentence, or within it) were not removed.
    reappearing = {footnote_text(r) for b in blocks if isinstance(b, (Paragraph, Heading)) for r in b.runs
                   if isinstance(r, Footnote) and (r.rev in ADDED or r.moved)}
    vanishing = {footnote_text(r) for b in blocks if isinstance(b, (Paragraph, Heading)) for r in b.runs
                 if isinstance(r, Footnote) and r.rev == "del"}
    # Footnotes that moved and were changed a little: the old one by the new one's text.
    removed_notes = {footnote_text(r): r for b in blocks if isinstance(b, (Paragraph, Heading)) for r in b.runs
                     if isinstance(r, Footnote) and r.rev == "del" and footnote_text(r) not in reappearing}
    changed_moves: dict[str, Footnote] = {}
    for b in blocks:
        for r in b.runs if isinstance(b, (Paragraph, Heading)) else []:
            text = footnote_text(r) if isinstance(r, Footnote) else ""
            if not (isinstance(r, Footnote) and r.rev == "ins" and not r.moved and text and text not in vanishing):
                continue
            best = max(((difflib.SequenceMatcher(None, text.split(), old.split(), autojunk=False).ratio(), old)
                        for old in removed_notes if old not in {footnote_text(n) for n in changed_moves.values()}),
                       default=(0.0, None))
            if best[0] >= 0.6:
                changed_moves[text] = removed_notes[best[1]]
    reappearing |= {footnote_text(note) for note in changed_moves.values()}

    # Where each removed footnote stood (its section in the old version), for "moved here from".
    footnote_places: dict[str, str] = {}
    section = ""
    for b in blocks:
        if isinstance(b, Heading) and b.mark_rev not in ADDED:
            title = "".join(r.text for r in b.runs if isinstance(r, Text) and r.rev not in ADDED).strip()
            number = "".join(r.text for r in b.number_runs or [] if isinstance(r, Text) and r.rev not in ADDED).strip() \
                if b.number_runs is not None else b.number
            section = f"section {number} “{title}”" if number else f"“{title}”"
        elif isinstance(b, Paragraph):
            for r in b.runs:
                if isinstance(r, Footnote) and r.rev == "del":
                    footnote_places.setdefault(footnote_text(r), section)

    def relocated(runs: list, here: str) -> list:
        """A footnote that only moved to another sentence: gone from the old place without a
        trace, blue (moved) at the new one, with a comment saying where it came from."""
        out_runs = []
        for r in runs:
            if isinstance(r, Footnote) and r.rev == "del" and footnote_text(r) in reappearing:
                continue
            note = None
            if isinstance(r, Footnote) and r.rev == "ins" and not r.moved and footnote_text(r) in changed_moves:
                # (moved and changed: the unchanged words blue, the new ones green)
                old_note = changed_moves[footnote_text(r)]
                origin = footnote_places.get(footnote_text(old_note), "")
                marked = diff_runs(runs_with_rev(old_note.runs, None), runs_with_rev(r.runs, None))
                r = replace(r, moved=True, runs=mark_deletion_sites(
                    [replace(x, rev="moveto") if isinstance(x, Text) and x.rev is None else x for x in marked]))
                note = CommentMark("This footnote was moved here from " + (
                    "elsewhere in this section" if not origin or origin == here else origin) + " and changed.")
            elif isinstance(r, Footnote) and r.rev == "ins" and not r.moved and footnote_text(r) in vanishing:
                r = replace(r, moved=True, runs=[replace(x, rev="moveto") if getattr(x, "rev", None) == "ins" else x
                                                 for x in r.runs])
                origin = footnote_places.get(footnote_text(r), "")
                note = CommentMark("This footnote was moved here from " + (
                    "elsewhere in this section." if not origin or origin == here else f"{origin}."))
            out_runs.append(r)
            if note is not None:
                out_runs.append(note)
        return out_runs

    def equation_next(position: int, rev: str) -> bool:
        """A paragraph mark changed only because a display equation inside a LaTeX paragraph was
        inserted or removed (Word paragraphs end at a display equation): no paragraph change."""
        following = blocks[position + 1] if position + 1 < len(blocks) else None
        return isinstance(following, Equation) and following.rev == rev

    out: list = []
    carry: list = []  # content of paragraphs whose paragraph mark was deleted
    pending: list[str] = []  # notes for the next paragraph (added paragraph breaks)
    removed_blocks: list = []  # removed headings and paragraphs since the last block shown
    relabelled: set = set()  # lists whose kind of label changed (noted at their first item)

    def flush_removed():
        # (Paragraphs removed before a removed heading belong to the section before it: their
        # own comment.)
        first = next((i for i, (kind, _) in enumerate(removed_blocks) if kind == "heading"), 0)
        for part in (removed_blocks[:first], removed_blocks[first:]):
            if part:
                out.append(RemovalNote(removal_note(part)))
        removed_blocks.clear()

    current = ""  # the section (of the new version) the block is in
    for position, block in enumerate(blocks):
        rev = getattr(block, "rev", None)
        if isinstance(block, Heading) and block.mark_rev not in REMOVED:
            title = "".join(r.text for r in block.runs if isinstance(r, Text) and r.rev not in REMOVED).strip()
            number = "".join(r.text for r in block.number_runs or [] if isinstance(r, Text)
                             and r.rev not in REMOVED).strip() if block.number_runs is not None else block.number
            current = f"section {number} “{title}”" if number else f"“{title}”"
        if isinstance(block, Equation) and rev == "del":
            if position in noted:
                note = "An equation was removed here" + (
                    f" (previously {equation_label(block)})." if block.number else ".")
                if carry:
                    # (It stood inside a paragraph, whose two parts are now one: the comment goes
                    # where it was.)
                    carry.append(CommentMark(note))
                else:
                    out.append(RemovalNote(note))
            continue
        if isinstance(block, Table) and rev is None and block.row_revs and "del" in block.row_revs:
            removed = sum(1 for row, row_rev in zip(block.rows, block.row_revs) if row_rev == "del" and not all(
                getattr(r, "moved", False) for c in row for r in c.runs if isinstance(r, Text) and r.text.strip()))
            if removed > 0:
                out.append(RemovalNote(f"{removed} row{'s' if removed > 1 else ''} of this table "
                                       f"{'were' if removed > 1 else 'was'} removed."))
            moved = sum(1 for row, row_rev in zip(block.rows, block.row_revs) if row_rev == "ins" and any(
                getattr(r, "moved", False) for c in row for r in c.runs if isinstance(r, Text) and r.text.strip()))
            if moved:
                out.append(RemovalNote(f"{moved} row{'s' if moved > 1 else ''} of this table "
                                       f"{'were' if moved > 1 else 'was'} moved (shown in blue)."))
        if isinstance(block, Table) and rev == "del" and block.replaced:
            continue  # the next table is its new version (changed columns)
        if isinstance(block, Table) and rev == "ins" and block.highlight_rows is not None:
            if block.removed_rows:
                out.append(RemovalNote(f"{block.removed_rows} row{'s' if block.removed_rows > 1 else ''} of this table "
                                       f"{'were' if block.removed_rows > 1 else 'was'} removed."))
            moved = sum(1 for row in block.highlight_rows if any(getattr(r, "moved", False) for c in row for r in c.runs))
            if moved:
                out.append(RemovalNote(f"{moved} row{'s' if moved > 1 else ''} of this table "
                                       f"{'were' if moved > 1 else 'was'} moved (shown in blue)."))
            block = copy(block)
            block.rows, block.row_revs, block.rev = block.highlight_rows, None, None
            block.caption_runs = block.highlight_caption
            block.rows = [[replace(c, runs=cell_deletions(c)) for c in row] for row in block.rows]
            out.append(block)
            continue
        if isinstance(block, Table) and rev is None:
            block = copy(block)
            if any(c.continues for row in block.rows for c in row):
                # (The new version's merged \multirow cell, without the old label kept for tracking.)
                def continued(c):
                    return replace(c, placeholder=True, runs=[], continues=False, rowspan=1, merge_start=False)
                block.rows = recompute_vertical_merges([[continued(c) if c.continues else c for c in row]
                                                        for row in block.rows])
            block.rows = [[replace(c, runs=cell_deletions(c, row_rev)) for c in row]
                          for row, row_rev in zip(block.rows, block.row_revs or [None] * len(block.rows))]
            if block.caption_runs:
                block.caption_runs = mark_deletion_sites(block.caption_runs, quote=True)
            for name in ("intro_runs", "outro_runs"):
                runs = getattr(block, name)
                if not runs:
                    continue
                if not any(isinstance(r, Text) and r.text.strip() and r.rev not in REMOVED for r in runs):
                    # (Text that went with the table - a panel heading, a note - removed: a comment
                    # on the caption, or on the table's first cell.)
                    quoted = quoted_text(" ".join(quoted_runs([r for r in runs if getattr(r, "rev", None) in REMOVED])
                                                  .split()), 15)
                    if quoted:
                        note = CommentMark(f"Text that went with this table was removed (“{quoted}”).")
                        if block.caption_runs:
                            block.caption_runs = [note] + list(block.caption_runs)
                        elif block.rows and block.rows[0]:
                            block.rows[0][0] = replace(block.rows[0][0], runs=[note] + list(block.rows[0][0].runs))
                    setattr(block, name, None)
                else:
                    setattr(block, name, mark_deletion_sites(runs, quote=True))
        if isinstance(block, Figure) and rev is None and block.caption_runs:
            block = copy(block)
            block.caption_runs = mark_deletion_sites(block.caption_runs, quote=True)
        if rev in REMOVED:
            if isinstance(block, (Figure, Table)) and rev == "del" and not block.move_note:
                kind = block.kind if isinstance(block, Figure) else "table"
                label = figure_label(block) if isinstance(block, Figure) else table_label(block)
                caption = " ".join((label + quoted_runs(block.caption)).split())
                if len(caption) > 90:
                    caption = caption[:90].rsplit(" ", 1)[0] + " …"
                out.append(RemovalNote(f"A {kind} was removed here" + (f" (previously “{caption}”)." if caption else ".")))
            continue
        if isinstance(block, (Paragraph, Heading)):
            def with_marks(runs):
                # A removed footnote leaves a comment where it was.
                out_runs = []
                for r in runs:
                    notes = []
                    if isinstance(r, Footnote) and r.rev is None:
                        # (A longer passage removed from the footnote: the comment goes on its mark.)
                        inner = mark_deletion_sites(r.runs, quote=True)
                        notes = [CommentMark(m.text.replace("Text was removed here", "Text was removed from this footnote"))
                                 for m in inner if isinstance(m, CommentMark)]
                        r = replace(r, runs=[m for m in inner if not isinstance(m, CommentMark)])
                    if isinstance(r, Footnote) and r.rev == "del" and footnote_text(r) not in reappearing:
                        words = quoted_runs(r.runs).split()
                        quoted = " ".join(words[:8]) + (" …" if len(words) > 8 else "")
                        out_runs.append(CommentMark(f"A footnote was removed here (“{quoted}”)." if quoted
                                                    else "A footnote was removed here."))
                    else:
                        out_runs.append(r)
                        out_runs.extend(notes)
                return out_runs

            if block.mark_rev == "del" and not any(
                    isinstance(r, Text) and r.text.strip() and getattr(r, "rev", None) not in REMOVED
                    for r in block.runs + (block.number_runs or [] if isinstance(block, Heading) else [])):
                # A removed paragraph or heading: one comment where a run of them was.
                old_text = " ".join(quoted_runs(block.runs).split())
                if isinstance(block, Heading):
                    number = "".join(r.text for r in block.number_runs or [] if isinstance(r, Text)) \
                        if block.number_runs is not None else block.number
                    removed_blocks.append(("heading", f"{number.strip()} {old_text}".strip()))
                elif old_text:
                    removed_blocks.append(("list item" if block.list_kind and block.list_item_start else
                                           "reference" if block.style.startswith("Bibliography") else
                                           "paragraph", old_text))
                continue
            flush_removed()
            if isinstance(block, Paragraph) and block.old_list is not None and block.list_id not in relabelled:
                old_kind, old_format = block.old_list
                kinds = {"bullet": "a bulleted", "number": "a numbered"}
                if old_kind in kinds and block.list_kind in kinds and old_kind != block.list_kind:
                    relabelled.add(block.list_id)
                    pending.append(f"This list was changed from {kinds[old_kind]} to {kinds[block.list_kind]} list.")
                elif old_kind == block.list_kind and old_format != block.list_format:
                    relabelled.add(block.list_id)
                    pending.append("The numbering style of this list was changed.")
            if isinstance(block, Heading):
                block = copy(block)
                block.runs = keep(mark_deletion_sites(block.runs))
                block.number_runs = keep(block.number_runs) if block.number_runs is not None else None
                content = keep(heading_runs(block))
            else:
                content = keep(with_marks(mark_deletion_sites(relocated(block.runs, current), quote=True)))
            unchanged = any(getattr(r, "rev", None) is None and isinstance(r, Text) and r.text.strip()
                            for r in block.runs)
            if block.mark_rev in REMOVED:
                carry.extend(content)
                if block.mark_rev == "del" and unchanged and content and not equation_next(position, "del"):
                    carry.append(CommentMark("Two paragraphs were joined here."))
                continue
            if block.mark_rev == "ins" and unchanged and isinstance(block, Paragraph) \
                    and not equation_next(position, "ins"):
                pending.append("A paragraph break was added after this text.")
            if isinstance(block, Heading):
                if carry:
                    block.runs = carry + content
                    block.number_runs = []
            else:
                block.runs = carry + content
            carry = []
            if isinstance(block, Paragraph) and not block.runs:
                continue
            if pending and isinstance(block, Paragraph):
                block = copy(block)
                block.notes = (block.notes or []) + pending
                pending = []
            out.append(block)
        else:
            flush_removed()
            if carry:
                out.append(Paragraph(carry))
                carry = []
            out.append(block)
    flush_removed()
    if carry:
        out.append(Paragraph(carry))
    return out


def cell_deletions(cell, row_rev: str | None = None) -> list:
    """Highlight mode: the deletion marks of a table cell, and a comment in a cell that was emptied
    (not in a \\multirow group's cells: its label merely moved to another row of the group)."""
    runs = cell.runs
    marked = mark_deletion_sites(runs)
    if cell.removed:  # (an emptied cell of a moved row)
        return marked + [CommentMark(f"The content of this cell was removed (“{quoted_text(cell.removed, 12)}”).")]
    deleted = "".join(r.text for r in runs if isinstance(r, Text) and r.rev == "del").strip()
    if deleted and row_rev is None and not (cell.placeholder or cell.rowspan > 1) and not any(
            isinstance(r, (Text, Math)) and getattr(r, "rev", None) not in REMOVED
            and (not isinstance(r, Text) or r.text.strip()) for r in runs):
        marked.append(CommentMark(f"The content of this cell was removed (“{quoted_text(deleted, 12)}”)."))
    return marked


def _comment_markup(element) -> bool:
    """A comment's range marks or reference run (or an empty run left for one)."""
    if element.tag in (qn("w:commentRangeStart"), qn("w:commentRangeEnd")):
        return True
    return element.tag == qn("w:r") and (element.find(qn("w:commentReference")) is not None or all(
        child.tag == qn("w:rPr") for child in element))


def paragraph_break_changed(block, nxt) -> bool:
    """BLOCK and the paragraph after it were split apart or joined (both keep text of both versions)."""
    def unchanged(runs) -> bool:
        return any(isinstance(r, Text) and r.rev is None and r.text.strip() for r in runs)

    return (isinstance(block, Paragraph) and block.mark_rev in ("ins", "del") and not block.move
            and isinstance(nxt, Paragraph) and not nxt.move and unchanged(block.runs) and unchanged(nxt.runs))


def removal_note(removed: list) -> str:
    """The comment on a run of removed headings and paragraphs: [(kind, text)]."""
    headings = [text for kind, text in removed if kind == "heading"]
    others = [(kind, text) for kind, text in removed if kind != "heading"]
    if headings:
        # (Only a heading: its text stays, or moved, under another heading.)
        names = " and ".join(f"“{quoted_text(title, 12)}”" for title in headings[:3])
        names += " …" if len(headings) > 3 else ""
        noun = ("section" if others else "heading") + ("s" if len(headings) > 1 else "")
        note = f"The {noun} {names} {'were' if len(headings) > 1 else 'was'} removed here"
        if others:
            note += f", with {len(others)} paragraph{'s' if len(others) > 1 else ''} of text"
        return note + "."
    kind = others[0][0] if all(k == others[0][0] for k, _ in others) else "paragraph"
    if len(others) == 1:
        return f"A {kind} was removed here (“{quoted_text(others[0][1])}”)."
    return f"{len(others)} {kind}s were removed here (the first: “{quoted_text(others[0][1], 15)}”)."


def accept_toc_deletions(entries: list) -> list:
    out = []
    carry: list = []
    for entry in entries:
        runs = carry + [r for r in entry.runs if getattr(r, "rev", None) != "del"]
        carry = []
        if entry.mark_rev == "del":
            carry = runs
            continue
        if runs:
            out.append(TocEntry(entry.level, runs, entry.bookmark, None))
    return out


def content_children(body) -> list:
    """Body children except the final section properties (new content goes before them)."""
    return [child for child in body if child.tag != qn("w:sectPr")]


class RemovalNote:
    """Highlight mode: a comment on the next block saying that a figure/table was removed."""

    def __init__(self, text: str):
        self.text = text


def versions_of(rev: str | None) -> set[str]:
    """The versions an element with revision mark REV exists in."""
    return {"new"} if rev in ADDED else {"old"} if rev in REMOVED else {"old", "new"}


def set_page_break_before(p) -> None:
    ppr = p.get_or_add_pPr()
    if ppr.find(qn("w:pageBreakBefore")) is None:
        ppr.append(OxmlElement("w:pageBreakBefore"))
        reorder_children(ppr, PPR_ORDER)


def mark_format_change(run_element) -> None:
    """Highlight mode: formatting-only changes get a dotted underline (no text was changed)."""
    rpr = run_element.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run_element.insert(0, rpr)
    remove_child(rpr, "w:u")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "dotted")
    rpr.append(underline)
    reorder_children(rpr, RPR_ORDER)


def heading_runs(block: Heading, unbold: bool = False, label: bool = True) -> list:
    if block.number_runs is not None:
        number = list(block.number_runs)
    elif block.number:
        number = [Text(block.number + " ", Style(), block.mark_rev if block.mark_rev in ("ins", "del") else None)]
    else:
        number = []
    runs = [r.with_style(bold=False) if unbold and isinstance(r, Text) and r.style.bold else r for r in block.runs]
    if block.label_prefix and number and label:
        # "Chapter 1" on a line of its own above the title (report/book classes). The space
        # after the number goes, in both versions: "Chapter {--2--}{++3++}".
        for version in (REMOVED, ADDED):
            for index in range(len(number) - 1, -1, -1):
                last = number[index]
                if isinstance(last, Text) and last.rev not in version:
                    if last.text != last.text.rstrip():
                        number[index] = replace(last, text=last.text.rstrip())
                    break
        # (The prefix belongs to both versions unless the whole number is new or deleted: a
        # renumbered "Chapter {--2--}{++3++}" keeps its "Chapter" and line break.)
        revs = {getattr(r, "rev", None) for r in number if isinstance(r, Text) and r.text.strip()}
        rev = revs.pop() if len(revs) == 1 else None
        if rev is None and block.old_label_prefix and block.old_label_prefix != block.label_prefix:
            prefix = [Text(block.old_label_prefix, Style(), "del"), Text(block.label_prefix, Style(), "ins")]
        else:
            prefix = [Text(block.label_prefix, Style(), rev)]
        number = prefix + number + [Break("line", rev)]
    return number + runs


def heading_fully_deleted(block: Heading) -> bool:
    runs = heading_runs(block)
    texts = [r for r in runs if isinstance(r, Text) and r.text.strip()]
    return bool(texts) and all(r.rev in REMOVED for r in texts) and block.mark_rev in REMOVED


def remap_rule_index(row_revs: list) -> dict:
    mapping = {}
    new_index = 0
    for index, rev in enumerate(row_revs):
        mapping[index] = new_index
        if rev != "del":
            new_index += 1
    mapping[len(row_revs)] = new_index
    return mapping


def remap_rules(rules: set, row_revs: list) -> set:
    """Row-rule indices after deleted rows were removed (highlight mode)."""
    mapping = {}
    new_index = 0
    for index, rev in enumerate(row_revs):
        mapping[index] = new_index
        if rev != "del":
            new_index += 1
    out = set()
    for rule in rules:
        out.add(mapping.get(rule, new_index))
    return out


def to_deleted(run_element) -> None:
    for t in run_element.iter(qn("w:t")):
        t.tag = qn("w:delText")
    for t in run_element.iter(qn("w:instrText")):
        t.tag = qn("w:delInstrText")


def add_text_elements(r, text: str) -> None:
    """Append w:t / w:tab / w:br children for text (tabs and newlines become elements)."""
    parts = re.split(r"(\t|\n)", text)
    for part in parts:
        if not part:
            continue
        if part == "\t":
            r.append(OxmlElement("w:tab"))
        elif part == "\n":
            r.append(OxmlElement("w:br"))
        else:
            t = OxmlElement("w:t")
            t.text = part
            if part != part.strip():
                t.set(qn("xml:space"), "preserve")
            r.append(t)


REVISION_TAGS = {"ins": "w:ins", "del": "w:del", "movefrom": "w:moveFrom", "moveto": "w:moveTo"}
RPR_ORDER = ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike",
             "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid",
             "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs",
             "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
             "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath", "w:rPrChange"]
PPR_ORDER = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
             "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
             "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
             "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
             "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
             "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"]


def run_properties(style: Style, size: float | None):
    rpr = OxmlElement("w:rPr")
    if style.mono:
        fonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), MONO_FONT)
        rpr.append(fonts)
    if style.bold:
        rpr.append(OxmlElement("w:b"))
    if style.italic:
        rpr.append(OxmlElement("w:i"))
    if style.smallcaps:
        rpr.append(OxmlElement("w:smallCaps"))
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(round(size * 2))))
        rpr.append(sz)
    if style.vert in ("super", "sub"):
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript" if style.vert == "super" else "subscript")
        rpr.append(va)
    if style.math:
        rpr.insert(len([c for c in rpr if c.tag in (qn("w:rFonts"), qn("w:b"), qn("w:i"), qn("w:smallCaps"))]),
                   OxmlElement("w:noProof"))  # (symbols such as F_{s,v} are not misspelt words)
    return rpr


def box_paragraph(paragraph, color: str) -> None:
    """Colored box around a paragraph (highlight mode: new/changed display equations)."""
    ppr = paragraph._p.get_or_add_pPr()
    remove_child(ppr, "w:pBdr")
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "12")
        element.set(qn("w:space"), "4")
        element.set(qn("w:color"), color)
        borders.append(element)
    ppr.append(borders)
    reorder_children(ppr, PPR_ORDER)


def frame_paragraph(paragraph, box: str) -> None:
    """The frame of a framed box (\\begin{framed}): a grey border, visible on light and dark pages;
    consecutive framed paragraphs share it."""
    ppr = paragraph._p.get_or_add_pPr()
    remove_child(ppr, "w:pBdr")
    borders = OxmlElement("w:pBdr")
    for edge in (("left",) if box == "left" else ("top", "left", "bottom", "right")):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "12" if box == "left" else "6")
        element.set(qn("w:space"), "4")
        element.set(qn("w:color"), "808080")
        borders.append(element)
    ppr.append(borders)
    reorder_children(ppr, PPR_ORDER)
    fmt = paragraph.paragraph_format
    fmt.left_indent = (fmt.left_indent or Pt(0)) + Pt(6)
    if box != "left":
        fmt.right_indent = (fmt.right_indent or Pt(0)) + Pt(6)


def add_shape_outline(sppr, color: str) -> None:
    """Colored frame around an inline picture (highlight mode: new/changed figures)."""
    ln = OxmlElement("a:ln")
    ln.set("w", str(int(Pt(3))))
    fill = OxmlElement("a:solidFill")
    clr = OxmlElement("a:srgbClr")
    clr.set("val", color)
    fill.append(clr)
    ln.append(fill)
    sppr.append(ln)


def insert_setting(settings, element) -> None:
    """Insert a child of w:settings at its schema position (trackRevisions comes early)."""
    order = ["w:writeProtection", "w:view", "w:zoom", "w:removePersonalInformation", "w:removeDateAndTime",
             "w:doNotDisplayPageBoundaries", "w:displayBackgroundShape", "w:printPostScriptOverText",
             "w:printFractionalCharacterWidth", "w:printFormsData", "w:embedTrueTypeFonts",
             "w:embedSystemFonts", "w:saveSubsetFonts", "w:saveFormsData", "w:mirrorMargins",
             "w:alignBordersAndEdges", "w:bordersDoNotSurroundHeader", "w:bordersDoNotSurroundFooter",
             "w:gutterAtTop", "w:hideSpellingErrors", "w:hideGrammaticalErrors", "w:activeWritingStyle",
             "w:proofState", "w:formsDesign", "w:attachedTemplate", "w:linkStyles",
             "w:stylePaneFormatFilter", "w:stylePaneSortMethod", "w:documentType", "w:mailMerge",
             "w:revisionView", "w:trackRevisions"]
    rank = {qn(tag): i for i, tag in enumerate(order)}
    target = rank.get(element.tag, len(order))
    if settings.find(element.tag) is not None:
        return
    for index, child in enumerate(settings):
        if rank.get(child.tag, len(order) + 1) > target:
            settings.insert(index, element)
            return
    settings.append(element)


# ---------------------------------------------------------------------- generic helpers


def is_page_break_paragraph(element) -> bool:
    """Whether a body element is a paragraph containing nothing but a page break."""
    if element.tag != qn("w:p"):
        return False
    if any(t.text for t in element.iter(qn("w:t"))) or element.find(".//" + qn("m:oMath")) is not None:
        return False
    if element.find(".//" + qn("w:drawing")) is not None:
        return False
    return any(br.get(qn("w:type")) == "page" for br in element.iter(qn("w:br")))


def heading_text(block: Heading) -> str:
    text = "".join(r.text for r in block.runs if isinstance(r, Text))
    return f"{block.number} {text}" if block.number else text


def make_auto(color) -> None:
    for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
        if color.get(qn(attr)) is not None:
            del color.attrib[qn(attr)]
    color.set(qn("w:val"), "auto")


def set_auto_color(rpr) -> None:
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
        reorder_children(rpr, RPR_ORDER)
    make_auto(color)


def set_font(style, name: str | None, size: float | None, bold: bool | None = None, italic: bool | None = None) -> None:
    font = style.font
    if name:
        font.name = name
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(attr)) is not None:
                del fonts.attrib[qn(attr)]
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), name)
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    # Automatic text color: black on white pages, white on dark (dark mode) pages.
    set_auto_color(style.element.get_or_add_rPr())


def remove_child(element, tag: str) -> None:
    if element is None:
        return
    for child in element.findall(qn(tag)):
        element.remove(child)


def set_table_borders(table, none: bool = True) -> None:
    tbl_pr = table._tbl.tblPr
    remove_child(tbl_pr, "w:tblStyle")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    remove_child(tbl_pr, "w:tblBorders")
    tbl_pr.append(borders)


def set_table_width(table, width_emu_or_in) -> None:
    tbl_pr = table._tbl.tblPr
    remove_child(tbl_pr, "w:tblW")
    width = OxmlElement("w:tblW")
    if isinstance(width_emu_or_in, float):
        twips = int(width_emu_or_in * 1440)
    else:
        twips = int(Emu(width_emu_or_in).inches * 1440)
    width.set(qn("w:w"), str(twips))
    width.set(qn("w:type"), "dxa")
    tbl_pr.append(width)


TBLPR_ORDER = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
               "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
               "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"]
TRPR_ORDER = ["w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore", "w:wAfter", "w:cantSplit",
              "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange"]


def reorder_children(element, order: list[str]) -> None:
    rank = {qn(tag): i for i, tag in enumerate(order)}
    children = sorted(list(element), key=lambda e: rank.get(e.tag, 99))
    for child in children:
        element.remove(child)
        element.append(child)


def set_grid_widths(table, widths_in: list[float]) -> None:
    grid = table._tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths_in):
        col.set(qn("w:w"), str(int(width * 1440)))


def set_table_layout_fixed(table) -> None:
    remove_child(table._tbl.tblPr, "w:tblLayout")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)


def _is_layout_table(tbl) -> bool:
    look = tbl.find(f"{qn('w:tblPr')}/{qn('w:tblLook')}")
    return look is not None and look.get(qn("w:val")) == "0600"


def mark_table_as_layout(table) -> None:
    """Accessibility hint: the table is used for layout only."""
    look = OxmlElement("w:tblLook")
    look.set(qn("w:val"), "0600")
    tbl_pr = table._tbl.tblPr
    remove_child(tbl_pr, "w:tblLook")
    tbl_pr.append(look)


def set_cell_border(cell, edge: str, size: int = 6, color: str = RULE_COLOR) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), "0")
    element.set(qn("w:color"), color)
    # Keep the schema order of tcBorders children (top, left/start, bottom, right/end).
    order = {qn("w:top"): 0, qn("w:left"): 1, qn("w:start"): 1, qn("w:bottom"): 2, qn("w:right"): 3, qn("w:end"): 3}
    children = sorted(borders, key=lambda e: order.get(e.tag, 9))
    for child in children:
        borders.remove(child)
        borders.append(child)


def set_cell_margins(cell, left: int, right: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    remove_child(tc_pr, "w:tcMar")
    margins = OxmlElement("w:tcMar")
    tc_pr.append(margins)
    for edge, value in (("left", left), ("right", right)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value * 2))
        element.set(qn("w:type"), "dxa")
        margins.append(element)


def set_vmerge(cell, value: str | None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    merge = OxmlElement("w:vMerge")
    if value:
        merge.set(qn("w:val"), value)
    remove_child(tc_pr, "w:vMerge")
    tc_pr.append(merge)
    reorder_tcpr(tc_pr)


def reorder_tcpr(tc_pr) -> None:
    order = ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
             "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"]
    reorder_children(tc_pr, order)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:cantSplit")
    tr_pr.append(element)


class _TcRef:
    """Minimal stand-in for a python-docx cell: the helpers below only need its w:tc."""

    def __init__(self, tc):
        self._tc = tc


def row_cells(row) -> list:
    """(cell, first grid column, column span) of every w:tc of a table row, in order."""
    out, col = [], 0
    for tc in row._tr.tc_lst:
        span = tc.grid_span
        out.append((_TcRef(tc), col, span))
        col += span
    return out


def fill_cell(cell, fill: str) -> None:
    """Cell background (\\rowcolor/\\cellcolor) with text that stays readable in dark mode."""
    from .colors import is_light

    tc_pr = cell._tc.get_or_add_tcPr()
    remove_child(tc_pr, "w:shd")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)
    reorder_tcpr(tc_pr)
    text_color = "1A1A1A" if is_light(fill) else "FFFFFF"
    for r in cell._tc.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        if rpr.find(qn("w:color")) is None and rpr.find(qn("w:shd")) is None:
            color = OxmlElement("w:color")
            color.set(qn("w:val"), text_color)
            rpr.append(color)
            reorder_children(rpr, RPR_ORDER)


def unique_cells(row) -> list:
    seen = []
    for cell in row.cells:
        if not any(cell._tc is s._tc for s in seen):
            seen.append(cell)
    return seen


def center_short_caption(paragraph, runs: list, tracked: bool) -> None:
    """LaTeX centres a caption that fits on one line."""
    def visible(rev):
        return "".join(r.text for r in runs if isinstance(r, Text) and getattr(r, "rev", None) not in rev)
    texts = [visible(REMOVED), visible(ADDED)] if tracked else [visible(REMOVED)]
    if all(text_width(text, 10.0) <= TEXT_WIDTH_IN * 72 * 0.97 for text in texts):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def paragraph_versions(p) -> set:
    """The versions ("old", "new") in which tracked paragraph P begins: where its paragraph mark
    exists, or where some of its content survives (and merges into the following paragraph)."""
    rpr = p.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
    out = set()
    for version, hidden in (("old", ("w:ins", "w:moveTo")), ("new", ("w:del", "w:moveFrom"))):
        hidden_tags = {qn(tag) for tag in hidden}
        if rpr is None or not any(child.tag in hidden_tags for child in rpr):
            out.add(version)
            continue
        for element in p.iter(qn("w:t"), qn("w:drawing"), "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"):
            ancestor = element.getparent()
            while ancestor is not None and ancestor is not p and ancestor.tag not in hidden_tags:
                ancestor = ancestor.getparent()
            if ancestor is p and (element.text or element.tag == qn("w:drawing")):
                out.add(version)
                break
    return out


def paragraph_style_id(p) -> str:
    style = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return style.get(qn("w:val"), "") if style is not None else ""


def text_width(text: str, size: float) -> float:
    """Rough width in points of text set in a Times-like font."""
    narrow, wide = set("[]()|.,;:'!ifjlrt "), set("MWmw@%")
    width = 0.0
    for char in text:
        if char in narrow:
            width += 0.3
        elif char in wide:
            width += 0.85
        elif char.isupper():
            width += 0.68
        else:
            width += 0.5
    return width * size


def estimate_column_widths(rows, ncols: int, size: float) -> list[float]:
    """Column widths in inches: the natural width of the longest cell, shrunk to the text width
    if necessary, but never narrower than the longest word of the column (no words broken)."""
    char_width = size * 0.52 / 72  # inches per character (proportional serif font, rough)
    mono_width = size * 0.62 / 72  # (a monospace font: identifiers such as max_summary_bytes)
    natural = [0.3] * ncols
    minimum = [0.3] * ncols
    for row in rows:
        col = 0
        for cell in row:
            if cell.colspan == 1 and col < ncols:
                texts = [(r.text, mono_width if r.style.mono else char_width) for r in cell.runs if isinstance(r, Text)]
                natural[col] = max(natural[col], sum(len(t) * w for t, w in texts) + 0.2)
                longest = max((len(word) * w for t, w in texts for word in t.split()), default=0)
                minimum[col] = max(minimum[col], longest + 0.2)
            col += cell.colspan
    if sum(natural) <= TEXT_WIDTH_IN:
        return natural
    if sum(minimum) >= TEXT_WIDTH_IN:
        return [w * TEXT_WIDTH_IN / sum(minimum) for w in minimum]
    spare = TEXT_WIDTH_IN - sum(minimum)
    stretch = [n - m for n, m in zip(natural, minimum)]
    total = sum(stretch) or 1.0
    return [m + spare * st / total for m, st in zip(minimum, stretch)]


def set_core_properties(doc, blocks: list, changes: ChangeOptions | None) -> None:
    """Replace python-docx's template metadata by the paper's title and sensible values."""
    props = doc.core_properties
    title = next((b for b in blocks if isinstance(b, Paragraph) and b.style == "Title"), None)
    props.title = re.sub(r"\s+", " ", "".join(
        (r.text if isinstance(r, Text) else " ") for r in title.runs
        if isinstance(r, (Text, Break)) and r.rev not in REMOVED)).strip() if title else ""
    props.subject = ""
    props.keywords = ""
    props.category = ""
    props.comments = "Converted from LaTeX with docxwright" + (
        "" if changes is None else
        f" (changes since {changes.baseline or 'an earlier version'} are "
        + ("highlighted)" if changes.mode == "highlight" else "tracked)"))
    author = changes.author if changes is not None else ""
    props.author = author
    props.last_modified_by = author
    props.revision = 1
    # Package properties are real UTC (python-docx writes naive datetimes with a "Z").
    now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    props.created = now
    props.modified = now
    # The template's blank page thumbnail and its "Microsoft Macintosh Word" application
    # properties do not describe this document.
    package = doc.part.package
    for rel_id, rel in list(package.rels.items()):
        if rel.reltype == RT.THUMBNAIL:
            del package.rels[rel_id]
        elif rel.reltype == RT.EXTENDED_PROPERTIES and not rel.is_external:
            rel.target_part._blob = (
                b"<?xml version='1.0' encoding='UTF-8' standalone='yes'?>\n"
                b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
                b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
                b"<Template>Normal.dotm</Template><Application>docxwright</Application><DocSecurity>0</DocSecurity>"
                b"</Properties>")


def separate_trailing_revisions(body, changes: ChangeOptions) -> None:
    """LibreOffice joins a tracked change at the end of a paragraph with a tracked change at the
    start of the next paragraph into one change, swallowing the (unchanged) paragraph mark in
    between; rejecting or accepting it would then join the following paragraph (a heading,
    say) onto this one. Changes whose time stamps differ by more than a minute are kept
    separate, so the change at the end of such a paragraph is dated two minutes earlier."""
    earlier = (changes.date - timedelta(minutes=2)).strftime("%Y-%m-%dT%H:%M:%SZ")
    wrappers = {qn("w:ins"), qn("w:del"), qn("w:moveTo"), qn("w:moveFrom")}
    skip = {qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:commentRangeStart"),
            qn("w:commentRangeEnd"), qn("w:moveFromRangeEnd"), qn("w:moveToRangeEnd")}
    for p in body.iter(qn("w:p")):
        ppr = p.find(qn("w:pPr"))
        rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
        if rpr is not None and any(rpr.find(tag) is not None for tag in wrappers):
            continue  # the paragraph mark itself is tracked
        for child in reversed(list(p)):
            if child.tag in skip:
                continue
            if child.tag == qn("w:r") and child.find(qn("w:commentReference")) is not None:
                continue
            if child.tag in wrappers:
                child.set(qn("w:date"), earlier)
            break


def _mark_of(p):
    rpr = p.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
    if rpr is None:
        return None
    for tag in ("w:ins", "w:del", "w:moveTo", "w:moveFrom"):
        element = rpr.find(qn(tag))
        if element is not None:
            return element
    return None


def _properties(p) -> bytes:
    from lxml import etree

    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        return b""
    ppr = deepcopy(ppr)
    for child in list(ppr):
        if child.tag in (qn("w:rPr"), qn("w:pPrChange"), qn("w:sectPr")):
            ppr.remove(child)
    return etree.tostring(ppr)


def _wholly(p, tags: tuple) -> bool:
    """Whether all text of paragraph P lies inside revision elements of the given TAGS."""
    tags = {qn(tag) for tag in tags}
    for t in p.iter(qn("w:t"), qn("w:delText"), qn("w:drawing")):
        ancestor = t.getparent()
        while ancestor is not None and ancestor is not p and ancestor.tag not in tags:
            ancestor = ancestor.getparent()
        if ancestor is p or ancestor is None:
            return False
    return True


def shift_paragraph_marks(body) -> None:
    """Paragraphs inserted, deleted or moved as a whole before a differently formatted
    paragraph, such as a heading: the tracked paragraph mark goes on the paragraph before them
    instead (an insertion or deletion), as Word records pressing Enter at its end. Rejecting
    (accepting) then removes the text and that mark, and the following heading keeps its style
    in LibreOffice too (which otherwise joins it to the removed paragraph). Only where the
    paragraph before has the same formatting; moved text keeps its move markup.

    Where the mark cannot be shifted (the paragraph before is formatted differently, or a
    table), an untracked empty paragraph of minimal height separates the tracked paragraphs
    from the heading: LibreOffice then joins them with that paragraph, and the heading keeps
    its style. (Word keeps the following paragraph's formatting anyway.)"""
    groups = {qn("w:ins"): ("w:ins", "w:moveTo"), qn("w:moveTo"): ("w:ins", "w:moveTo"),
              qn("w:del"): ("w:del", "w:moveFrom"), qn("w:moveFrom"): ("w:del", "w:moveFrom")}
    children = [c for c in body if c.tag in (qn("w:p"), qn("w:tbl"))]
    ids = [int(e.get(qn("w:id"))) for e in body.iter() if (e.get(qn("w:id")) or "").isdigit()]
    next_id = max(ids, default=0) + 1

    for index, p in enumerate(children):
        mark = _mark_of(p) if p.tag == qn("w:p") else None
        if mark is None or not _wholly(p, groups[mark.tag]):
            continue
        tags = groups[mark.tag]
        following = children[index + 1] if index + 1 < len(children) else None
        if following is not None and following.tag == qn("w:p") and _mark_of(following) is not None \
                and _mark_of(following).tag in {qn(t) for t in tags} and _wholly(following, tags):
            continue  # (not the last of the inserted/deleted paragraphs: that one is handled)
        if following is None or following.tag != qn("w:p") or _properties(following) == _properties(p):
            continue
        start = index
        while start > 0 and children[start - 1].tag == qn("w:p"):
            previous = _mark_of(children[start - 1])
            if previous is None or previous.tag not in {qn(t) for t in tags} or not _wholly(children[start - 1], tags):
                break
            start -= 1
        before = children[start - 1] if start > 0 else None
        if before is None or before.tag != qn("w:p") or _mark_of(before) is not None \
                or _properties(before) != _properties(p):
            if not _merges_into(children, start, tags):
                p.addnext(_spacer_paragraph())
                # (LibreOffice gives what is left of the joined paragraphs the formatting of the
                # first of them - a figure's page break, a heading style -: a tracked empty
                # paragraph of minimal height leads them.)
                lead = _spacer_paragraph()
                lead_mark = OxmlElement(tags[0])
                for attr, value in mark.attrib.items():
                    lead_mark.set(attr, value)
                lead_mark.set(qn("w:id"), str(next_id))
                next_id += 1
                lead.find(qn("w:pPr")).find(qn("w:rPr")).insert(0, lead_mark)
                children[start].addprevious(lead)
            continue
        mark.getparent().remove(mark)
        shifted = OxmlElement(tags[0])  # (a moved paragraph's mark becomes a plain insertion/deletion)
        for attr, value in mark.attrib.items():
            shifted.set(attr, value)
        ppr = before.get_or_add_pPr()
        rpr = ppr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ppr.append(rpr)
            reorder_children(ppr, PPR_ORDER)
        rpr.insert(0, shifted)


def merge_replaced_paragraphs(body, revision_element=None) -> None:
    """A wholly deleted paragraph directly followed by a wholly inserted one with the same
    formatting (a rewritten paragraph, a list item replaced by another): one paragraph holding
    the deleted and the inserted text, with an untracked paragraph mark. Accepting or rejecting
    gives exactly the same, but LibreOffice has no tracked paragraph marks to mishandle (it
    would otherwise give the next paragraph, or an empty bullet, the removed one's style)."""
    removed, added = ("w:del", "w:moveFrom"), ("w:ins", "w:moveTo")
    children = [c for c in body if c.tag in (qn("w:p"), qn("w:tbl"))]
    for first, second in zip(children, children[1:]):
        if first.tag != qn("w:p") or second.tag != qn("w:p") or first.getparent() is None:
            continue
        mark_first, mark_second = _mark_of(first), _mark_of(second)
        if mark_first is None or mark_second is None or mark_first.tag not in {qn(t) for t in removed} \
                or mark_second.tag not in {qn(t) for t in added}:
            continue
        if not _wholly(first, removed) or not _wholly(second, added) or _properties(first) != _properties(second):
            continue
        if any(True for _ in first.iter(qn("w:drawing"), qn("w:sectPr"), qn("m:oMath"))) or any(
                True for _ in second.iter(qn("w:drawing"), qn("w:sectPr"), qn("m:oMath"))):
            continue  # (pictures, and formulas, whose changes LibreOffice cannot apply)
        ppr = second.find(qn("w:pPr"))
        position = list(second).index(ppr) + 1 if ppr is not None else 0
        for child in [c for c in first if c.tag != qn("w:pPr")]:
            second.insert(position, child)
            position += 1
        if revision_element is not None:
            # A deleted space between the old and the new text ("step. A ..." rather than
            # "step.A ..." in the markup view).
            space = revision_element("w:del")
            run = OxmlElement("w:r")
            text = OxmlElement("w:delText")
            text.set(qn("xml:space"), "preserve")
            text.text = " "
            run.append(text)
            space.append(run)
            second.insert(position, space)
        rpr = mark_second.getparent()
        rpr.remove(mark_second)
        if not len(rpr):
            rpr.getparent().remove(rpr)
        body.remove(first)


def plain_move_marks(body) -> None:
    """The paragraph mark of a moved paragraph as a plain insertion/deletion (its text stays a
    move): LibreOffice leaves an emptied paragraph - an empty bullet of a moved list item -
    when it rejects a paragraph whose mark is a move. Word accepts and rejects both alike."""
    for tag, plain in (("w:moveTo", "w:ins"), ("w:moveFrom", "w:del")):
        for mark in body.iter(qn(tag)):
            parent = mark.getparent()
            if parent is not None and parent.tag == qn("w:rPr") and parent.getparent() is not None \
                    and parent.getparent().tag == qn("w:pPr"):
                mark.tag = qn(plain)


def _merges_into(children: list, start: int, tags: tuple) -> bool:
    """Whether an earlier paragraph, whose mark is tracked like the paragraphs from START on,
    joins them (and the paragraph after them) when the change is rejected or accepted: a
    paragraph split around an inserted table or equation, say. Tables and paragraphs that
    vanish along with them are skipped."""
    tracked = {qn(t) for t in tags}
    index = start - 1
    while index >= 0:
        child = children[index]
        if child.tag == qn("w:tbl"):
            rows = child.findall(qn("w:tr"))
            if not rows or not all(tr.find(qn("w:trPr")) is not None and any(
                    tr.find(qn("w:trPr")).find(tag) is not None for tag in tracked) for tr in rows):
                return False
        else:
            mark = _mark_of(child)
            if mark is None or mark.tag not in tracked:
                return False
            if not _wholly(child, tags):
                return True
        index -= 1
    return False


def _spacer_paragraph():
    spacer = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    for attr, value in (("w:before", "0"), ("w:after", "0"), ("w:line", "20"), ("w:lineRule", "exact")):
        spacing.set(qn(attr), value)
    ppr.append(spacing)
    rpr = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "2")
    rpr.append(size)
    ppr.append(rpr)
    spacer.append(ppr)
    return spacer


LANGUAGE_TAGS = {"english": "en-US", "german": "de-DE", "french": "fr-FR", "spanish": "es-ES", "italian": "it-IT"}


def set_document_language(doc, language: str) -> None:
    """Word's proofing and hyphenation language for the whole document (docDefaults)."""
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is None:
        return
    rpr = defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}")
    if rpr is None:
        return
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
        reorder_children(rpr, RPR_ORDER)
    lang.set(qn("w:val"), LANGUAGE_TAGS.get(language, "en-US"))


def set_run_language(r, tag: str) -> None:
    """A run's proofing language (also in the formatting it had before a tracked format change,
    which would otherwise show the language as changed)."""
    rpr = r.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        r.insert(0, rpr)
    targets = [rpr] + [old for change in rpr.findall(qn("w:rPrChange")) for old in change.findall(qn("w:rPr"))]
    for target in targets:
        lang = target.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            target.append(lang)
            reorder_children(target, RPR_ORDER)
        lang.set(qn("w:val"), tag)


def write_docx(document, path: Path, changes: ChangeOptions | None = None) -> None:
    writer = DocxWriter(changes)
    writer.language = getattr(document, "language", "english")
    set_document_language(writer.doc, writer.language)
    writer.write(document.blocks)
    if writer.tracked:
        merge_replaced_paragraphs(writer.doc.element.body, writer.revision_element)
        shift_paragraph_marks(writer.doc.element.body)
        plain_move_marks(writer.doc.element.body)
        separate_trailing_revisions(writer.doc.element.body, changes)
    set_core_properties(writer.doc, document.blocks, changes)
    body = writer.doc.element.body
    for cell_fix in body.iter(qn("w:tcPr")):
        reorder_tcpr(cell_fix)
    for tr_pr in body.iter(qn("w:trPr")):
        reorder_children(tr_pr, TRPR_ORDER)
    for tbl_pr in body.iter(qn("w:tblPr")):
        look = tbl_pr.find(qn("w:tblLook"))
        if look is not None:
            # Keep only the ECMA-376 1st edition attribute (understood by every Word version).
            for attr in ("w:firstRow", "w:lastRow", "w:firstColumn", "w:lastColumn", "w:noHBand", "w:noVBand"):
                if look.get(qn(attr)) is not None:
                    del look.attrib[qn(attr)]
        reorder_children(tbl_pr, TBLPR_ORDER)
    for wrapper in body.iter(qn("w:ins"), qn("w:del")):
        wrapper.attrib.pop("_merge", None)
    renumber_comments(writer.doc, writer.footnotes)
    writer.save(path)


COMMENT_TAGS = ("w:commentRangeStart", "w:commentRangeEnd", "w:commentReference")


def renumber_comments(doc, footnotes) -> None:
    """Give comments ids above all other annotation ids (python-docx numbers them from 0).

    Ids must be unique across comments, revisions and bookmarks, and LibreOffice 7.4 shows the
    wrong footnote texts when a comment between two footnote references has a footnote's id."""
    body = doc.element.body
    anchors = list(body.iter(*(qn(tag) for tag in COMMENT_TAGS)))
    if not anchors:
        return
    skip = {qn(tag) for tag in COMMENT_TAGS}
    used = [0]
    for root in (body, footnotes):
        if root is None:
            continue
        for element in root.iter():
            value = element.get(qn("w:id"))
            if element.tag not in skip and value and value.lstrip("-").isdigit():
                used.append(int(value))
    offset = max(used) + 1
    for element in anchors:
        element.set(qn("w:id"), str(int(element.get(qn("w:id"))) + offset))
    for comment in doc.part._comments_part.element.iter(qn("w:comment")):
        comment.set(qn("w:id"), str(int(comment.get(qn("w:id"))) + offset))
